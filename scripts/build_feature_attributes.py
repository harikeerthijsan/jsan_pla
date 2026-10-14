"""Build api/app/feature_attributes.py from the "Feature and required attribute fields" workbook.

The sheet has one block per feature: a header cell (Pole Base, attachments, CROSSARM, ...) over two columns, field
names in the left column and their dropdown items below/next to them in the right column. A cell such as
"INCOMPLETE / YES" lists two items; a cell containing "MANUAL ENTRY" means the field is typed in, not picked.

Usage (from the repo root):  python scripts/build_feature_attributes.py "<path to workbook.xlsx>"
"""
from __future__ import annotations

import pprint
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "api" / "app" / "feature_attributes.py"

# Workbook block header -> application annotation group(s). Attachment fields split by prefix.
BLOCKS = {
    "pole base": ("poles", ["Pole_Base"]),
    "crossarm": ("crossarms", None),
    "equipment": ("equipment", None),
    "anchors": ("anchors", None),
    "guys": ("guys", None),
    "span_guys": ("span_guys", None),
    "sidewalk_braces": ("sidewalk_braces", None),
}
def clean(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def read_blocks(path: Path) -> dict[str, list[dict]]:
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    groups: dict[str, list[dict]] = {}
    for col in range(1, ws.max_column + 1):
        header = clean(ws.cell(1, col).value).lower()
        if not header:
            continue
        fields, current = [], None
        for row in range(2, ws.max_row + 1):
            name, value = clean(ws.cell(row, col).value), clean(ws.cell(row, col + 1).value)
            if name:
                # The field name is shown exactly as written in the workbook.
                current = {"key": name, "label": name, "options": [], "manual": False}
                fields.append(current)
            if current is None or not value:
                continue
            if "MANUAL ENTRY" in value.upper():
                current["manual"] = True
                continue
            for item in [v.strip() for v in value.split(" / ")] if " / " in value and value.isupper() else [value]:
                if item and item not in current["options"]:
                    current["options"].append(item)
        if header == "attachments":
            for field in fields:
                family = "attachments_comm" if field["key"].startswith("comm_") else "attachments_util"
                groups.setdefault(family, []).append(field)
            continue
        if header not in BLOCKS:
            raise SystemExit(f"Unknown block header in column {col}: {ws.cell(1, col).value!r}")
        family, features = BLOCKS[header]
        for field in fields:
            field["features"] = features
        groups.setdefault(family, []).extend(fields)
    return groups


# Fields that are not in the workbook but are required in the app, appended after the workbook's own fields.
# "source": "anchors" fills the dropdown with the anchor points already saved for the pole (anc_1, anc_2, ...).
EXTRA_FIELDS = {
    # COLLECTION lut_enums: first_pole / last_pole INCOMPLETE -> prev_inc / next_inc lat, lon, elev. Shown on Pole_Base only when
    # that field is INCOMPLETE; "pick" fills all three from a click in the LiDAR view.
    "poles": [
        {"key": f"{side}_{part}", "label": f"{side}_{part}", "options": [], "manual": True, "features": ["Pole_Base"],
         "pick": side, "show_when": {flag: "INCOMPLETE"}}
        for side, flag in (("prev_inc", "first_pole"), ("next_inc", "last_pole")) for part in ("lat", "lon", "elev_ft")
    ],
    # util_N_arm_no is a COLLECTION workbook column; the attribute workbook has no utility arm field, so it is added here.
    "attachments_util": [
        {"key": "util_arm_no", "label": "util_arm_no", "options": [], "manual": False, "source": "crossarms", "features": None},
    ],
    "guys": [
        {"key": "guy_anchor_no", "label": "guy_anchor_no", "options": [], "manual": False, "source": "anchors", "features": None},
    ],
    "sidewalk_braces": [
        {"key": "swb_anchor_no", "label": "swb_anchor_no", "options": [], "manual": False, "source": "anchors", "features": None},
        {"key": "swb_length_ft", "label": "swb_length_ft", "options": [], "manual": True, "features": None},
        {"key": "swb_guy_nos", "label": "swb_guy_nos", "options": [], "manual": True, "features": None},
    ],
}


# Fields from the workbook whose dropdown the app fills in: the arm the attachment hangs from (the crossarms saved
# on the pole, or "pole"), the same choices as the point's Parent.
FIELD_OVERRIDES = {
    ("attachments_comm", "comm_arm_no"): {"manual": False, "source": "crossarms", "options": []},
}


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__)
        return 2
    groups = read_blocks(Path(argv[0]))
    for (family, key), change in FIELD_OVERRIDES.items():
        for field in groups.get(family, []):
            if field["key"] == key:
                field.update(change)
    for family, extra in EXTRA_FIELDS.items():
        present = {f["key"] for f in groups.setdefault(family, [])}
        groups[family].extend(f for f in extra if f["key"] not in present)
    body = pprint.pformat(groups, width=120, sort_dicts=False)
    OUT.write_text(
        '"""Attribute fields and dropdown items per production annotation group.\n\n'
        "Generated by scripts/build_feature_attributes.py from the \"Feature and required attribute fields\" workbook;\n"
        "re-run the script when the workbook changes rather than editing by hand. A field with \"manual\": True is typed\n"
        "in; otherwise the user picks an option or chooses Other and types a value. \"features\" limits a field to those\n"
        'point types (None = every point in the group)."""\n\n'
        f"FEATURE_ATTRIBUTES = {body}\n",
        encoding="utf-8",
    )
    total = sum(len(f) for f in groups.values())
    print(f"wrote {OUT.relative_to(ROOT)}: {total} fields across {len(groups)} groups")
    for family, fields in groups.items():
        print(f"  {family}: " + ", ".join(f"{f['key']}({'manual' if f['manual'] else len(f['options'])})" for f in fields))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
