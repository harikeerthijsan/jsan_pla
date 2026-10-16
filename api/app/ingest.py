import json, math, re, uuid
from collections import Counter, defaultdict
from openpyxl import load_workbook
from pyproj import Transformer
from .source_schema import detect_pole_sheet, field_key, has_location, header_map, normalize_header

EXPECTED_SHEETS=['poles','attachments','crossarms','equipment','anchors','guys','sidewalk_braces','span_guys']

def val(v):
    if v is None: return None
    if isinstance(v,str):
        v=v.strip()
        return v if v else None
    return v

def num(v):
    try:
        if v is None or v=='': return None
        if isinstance(v,str): v=v.strip().lstrip(',').strip()
        return float(v)
    except Exception: return None

def as_int(v):
    n=num(v)
    return int(n) if n is not None and float(n).is_integer() else None

def rows_by_header(ws, *, require_identity=False):
    headers=[val(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1),())]
    canonical=header_map(headers)
    for row_number,row in enumerate(ws.iter_rows(min_row=2, values_only=True),start=2):
        source={str(headers[i]):val(row[i]) for i in range(min(len(headers),len(row))) if headers[i]}
        if not any(value is not None for value in source.values()): continue
        d={field_key(headers[i]):val(row[i]) for i in range(min(len(headers),len(row))) if headers[i]}
        for name,index in canonical.items():
            if index<len(row): d[name]=val(row[index])
        d['__source__']=source
        if require_identity and as_int(d.get('internal_id')) is None:
            if d.get('internal_id') is not None:
                raise ValueError(f"Pole catalogue worksheet {ws.title!r} row {row_number} has an invalid Internal ID")
            if d.get('pole_number') is None:
                continue  # a note or stray cell, not a pole
            # Internal ID left blank: load the pole with a provisional id (negative, so it can never clash with a real
            # one, and stable for the same row); it is set later in the Production view, which writes it to the Excel.
            d['internal_id']=-row_number; d['__provisional__']=True
        if as_int(d.get('internal_id')) is not None:
            yield d

# Every rule the QC run can apply. "needs" explains what the workbook must contain for the rule to run, so a run on a
# differently shaped workbook reports the rule as not applicable instead of silently passing it.
RULES={
    'PLA-R002':('PLA','Pole number is unique','Pole Number column'),
    'PLA-R003':('PLA','Pole has complete base/top 3D geometry','latitude, longitude, bottom_elev_ft and top_elev_ft columns on the pole catalogue'),
    'PLA-R006':('PLA','Other-pole connections are consistent','other_N_id / other_N_type columns on the pole catalogue'),
    'PLA-R014':('PLA','Grade of construction is filled','a grade_of_construction column on the pole catalogue'),
    'PLA-R020':('PLA','Equipment, anchor and crossarm XYZ is complete','equipment / anchors / crossarms sheets with eq_N_, anc_N_ or arm_N_ columns'),
    'PLA-R021':('PLA','Equipment with a key has an owner','an equipment sheet with eq_N_key and eq_N_owner columns'),
    'PLA-R010':('PLA','Guys reference an existing anchor','a guys sheet with guy_N_anchor_no columns'),
    'PLA-R012':('PLA','Sidewalk braces list their guys','a sidewalk_braces sheet with swb_N_guy_nos columns'),
    'PLA-R030':('PLA','Span guys reference a defined Other connection','a span_guys sheet with sgy_N_runs_to columns'),
    'GEN-001':('GENERIC','Pole has a usable location','latitude/longitude or X/Y columns'),
    'GEN-002':('GENERIC','Coordinates are valid numbers in range','latitude/longitude or X/Y columns'),
    'GEN-003':('GENERIC','Pole lies inside the LiDAR coverage','processed LiDAR tiles'),
    'GEN-004':('GENERIC','Pole has a matching GeoJSON point','an uploaded GeoJSON file'),
    'GEN-005':('GENERIC','GeoJSON point agrees with the workbook location','an uploaded GeoJSON file'),
    'GEN-006':('GENERIC','Only one GeoJSON point matches the pole','an uploaded GeoJSON file'),
}

def rule_status(rule_id,applicable,reason=None):
    scope,title,needs=RULES[rule_id]
    return {'rule_id':rule_id,'scope':scope,'title':title,'applicable':bool(applicable),
            'reason':None if applicable else (reason or f'Not applicable: needs {needs}')}

def make_finding(rule,severity,pid,pnum,sheet,field,message,actual=None,expected=None,related=None):
    return {'id':f'F-{uuid.uuid4().hex[:10]}','rule_id':rule,'severity':severity,
        'internal_id':pid,'pole_number':pnum,'sheet':sheet,'field':field,'message':message,
        'actual':None if actual is None else str(actual),'expected':expected,'related_poles':related or [],'status':'OPEN'}

def sheet_fields(ws):
    """Normalised column keys of a worksheet header (plus canonical identity/location names)."""
    headers=[val(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1),())]
    return {field_key(h) for h in headers if h}|set(header_map(headers))

def _has(fields,pattern):
    return any(re.fullmatch(pattern,name) for name in fields)

def project_transformer(crs):
    return Transformer.from_crs('EPSG:4326',crs,always_xy=True)

def point_feature(internal_id, feature_type, xyz, **props):
    p={'feature_type':feature_type,'internal_id':internal_id,**props}
    return {'type':'Feature','properties':p,'geometry':{'type':'Point','coordinates':[float(x) for x in xyz]}}

def line_feature(internal_id, feature_type, a,b, **props):
    p={'feature_type':feature_type,'internal_id':internal_id,**props}
    return {'type':'Feature','properties':p,'geometry':{'type':'LineString','coordinates':[[float(x) for x in a],[float(x) for x in b]]}}

def parse_workbook(path:str, crs='EPSG:6424'):
    wb=load_workbook(path, read_only=True, data_only=True)
    pole_sheet,_=detect_pole_sheet(wb,require_location=True)
    tf=project_transformer(crs)
    sheets={s:[] for s in EXPECTED_SHEETS}; fields={s:set() for s in EXPECTED_SHEETS}
    sheets['poles']=list(rows_by_header(pole_sheet,require_identity=True)); fields['poles']=sheet_fields(pole_sheet)
    pole_headers=[val(c.value) for c in next(pole_sheet.iter_rows(min_row=1,max_row=1),())]
    column_name={name:str(pole_headers[index]) for name,index in header_map(pole_headers).items()}
    for worksheet in wb.worksheets:
        key=field_key(worksheet.title)
        if key in sheets and key!='poles': sheets[key]=list(rows_by_header(worksheet)); fields[key]=sheet_fields(worksheet)
    if not sheets['poles']: raise ValueError(f"Pole catalogue worksheet {pole_sheet.title!r} has no valid data rows")
    poles=[]; features=[]; findings=[]; byid={}

    def addfinding(rule,severity,pid,pnum,sheet,field,message,actual=None,expected=None,related=None):
        findings.append(make_finding(rule,severity,pid,pnum,sheet,field,message,actual,expected,related))

    pf=fields['poles']
    applies={
        'PLA-R002':True,
        'PLA-R003':{'latitude','longitude','bottom_elev_ft','top_elev_ft'}<=pf,
        'PLA-R006':_has(pf,r'other_\d+_(id|type)'),
        'PLA-R014':'grade_of_construction' in pf,
        'PLA-R020':any(_has(fields[sheet],rf'{prefix}_\d+_.+') for sheet,prefix in (('equipment','eq'),('anchors','anc'),('crossarms','arm'))),
        'PLA-R021':_has(fields['equipment'],r'eq_\d+_key') and _has(fields['equipment'],r'eq_\d+_owner'),
        'PLA-R010':_has(fields['guys'],r'guy_\d+_anchor_no'),
        'PLA-R012':_has(fields['sidewalk_braces'],r'swb_\d+_guy_nos'),
        'PLA-R030':_has(fields['span_guys'],r'sgy_\d+_runs_to'),
        'GEN-001':True,'GEN-002':True,
    }

    def coordinate_problems(r):
        """GEN-002: a coordinate cell that has a value which is not a usable number (or is out of range)."""
        problems=[]
        for name,low,high in (('latitude',-90,90),('longitude',-180,180),('x',None,None),('y',None,None)):
            raw=r.get(name)
            if raw is None: continue
            value=num(raw); column=column_name.get(name,name)
            if value is None or not math.isfinite(value):
                problems.append((column,raw,f'{column} {raw!r} is not a number'))
            elif low is not None and not low<=value<=high:
                problems.append((column,raw,f'{column} {value} is outside {low}..{high}'))
        lat,lon=num(r.get('latitude')),num(r.get('longitude'))
        if lat is not None and lon is not None and abs(lat)>90 and abs(lon)<=90:
            problems.append((f"{column_name.get('latitude','latitude')}/{column_name.get('longitude','longitude')}",f'{lat}, {lon}','Latitude and longitude look swapped'))
        return problems

    for r in sheets['poles']:
        pid=as_int(r.get('internal_id')); pnum=val(r.get('pole_number'))
        lat=num(r.get('latitude')); lon=num(r.get('longitude')); z0=num(r.get('bottom_elev_ft')); z1=num(r.get('top_elev_ft'))
        source_x=num(r.get('x')); source_y=num(r.get('y')); source_z=num(r.get('z'))
        top_lat=num(r.get('top_lat')) or lat; top_lon=num(r.get('top_lon')) or lon
        remarks=val(r.get('remarks'))
        manifest={k:v for k,v in r['__source__'].items() if v is not None}
        if not pnum: raise ValueError(f"Pole catalogue row for Internal ID {pid} has no Pole Number")
        if pid in byid: raise ValueError(f"Pole catalogue has duplicate Internal ID {pid}")
        p={'internal_id':pid,'pole_number':pnum,'corrected_lat':lat,'corrected_lon':lon,'source_x':source_x,'source_y':source_y,'source_z':source_z,
           'bottom_elev_ft':z0,'top_elev_ft':z1,'remarks':remarks,'manifest':manifest,'provisional_internal_id':bool(r.get('__provisional__'))}
        poles.append(p); byid[pid]=p
        bad=coordinate_problems(r)
        for column,raw,message in bad:
            addfinding('GEN-002','FAIL',pid,pnum,pole_sheet.title,column,message+'.',raw,'a number in range')
        if bad:  # unusable values must not place the pole anywhere
            lat=lon=source_x=source_y=None; p.update(corrected_lat=None,corrected_lon=None,source_x=None,source_y=None)
            top_lat=num(r.get('top_lat')); top_lon=num(r.get('top_lon'))
        has_xy=source_x is not None and source_y is not None
        complete_latlon=None not in (lat,lon,z0,top_lat,top_lon,z1)
        if not has_xy and (lat is None or lon is None) and not bad and not applies['PLA-R003']:
            addfinding('GEN-001','FAIL',pid,pnum,pole_sheet.title,'location','Pole has no usable location (latitude/longitude or X/Y), so it cannot be found in the LiDAR.',
                       f'lat={lat}, lon={lon}, x={source_x}, y={source_y}','latitude/longitude or X/Y')
        if has_xy:
            # Workbook X/Y is already in the project CRS; keep base/top elevations when the sheet has them.
            if z0 is not None and z1 is not None:
                features.append(line_feature(pid,'pole',[source_x,source_y,z0],[source_x,source_y,z1],pole_number=pnum,bottom_elev_ft=z0,top_elev_ft=z1,source='workbook_xyz'))
            elif source_z is not None: features.append(point_feature(pid,'pole_location',[source_x,source_y,source_z],pole_number=pnum,source='workbook_xyz'))
        elif not complete_latlon and applies['PLA-R003'] and not bad:
            sev='UNVERIFIABLE' if remarks and 'not visible' in remarks.lower() else 'FAIL'
            addfinding('PLA-R003',sev,pid,pnum,'poles','latitude/longitude/elevation','Pole lacks complete measurable 3D geometry.',
                       actual=f'lat={lat}, lon={lon}, bottom={z0}, top={z1}',expected='complete base/top geometry')
        elif complete_latlon:
            x0,y0=tf.transform(lon,lat); x1,y1=tf.transform(top_lon,top_lat)
            features.append(line_feature(pid,'pole',[x0,y0,z0],[x1,y1,z1],pole_number=pnum,bottom_elev_ft=z0,top_elev_ft=z1))
        if any(field_key(key)=='grade_of_construction' for key in r['__source__']) and not val(r.get('grade_of_construction')):
            addfinding('PLA-R014','REVIEW',pid,pnum,'poles','grade_of_construction','Grade of construction is blank.',None,'verify grade or accepted blank')
        for n in range(1,6):
            typ=val(r.get(f'other_{n}_type')); oid=as_int(r.get(f'other_{n}_id'))
            olat=num(r.get(f'other_{n}_lat')); olon=num(r.get(f'other_{n}_lon')); oe=num(r.get(f'other_{n}_elev_ft'))
            if oid==pid:
                addfinding('PLA-R006','FAIL',pid,pnum,'poles',f'other_{n}_id',f'Other connection self-references internal_id {pid}.',oid,'different internal_id')
            if oid and (olat is not None or olon is not None or oe is not None):
                addfinding('PLA-R006','FAIL',pid,pnum,'poles',f'other_{n}_xyz',f'In-scope Other {n} has internal_id and XYZ; in-scope connections should use the pole reference.',
                           f'id={oid}, xyz={olat},{olon},{oe}','ID reference without duplicate XYZ')

    # duplicate pole numbers -> REVIEW
    nums=defaultdict(list)
    for p in poles:
        if p['pole_number']: nums[p['pole_number']].append(p['internal_id'])
    for pnum,ids in nums.items():
        if len(ids)>1:
            for pid in ids:
                addfinding('PLA-R002','REVIEW',pid,pnum,'poles','pole_number',f'pole_number {pnum} appears on multiple internal_ids {ids}',pnum,'unique unless verified',[x for x in ids if x!=pid])

    # generic point layers
    specs=[
      ('attachments','comm','attachment',9,{'owner':'owner','size':'size','support':'support','arm_no':'arm_no','lat':'lat','lon':'lon','elev_ft':'elev_ft'}),
      ('attachments','util','attachment',21,{'usage':'usage','owner':'owner','size':'size','support':'support','arm_no':'arm_no','lat':'lat','lon':'lon','elev_ft':'elev_ft'}),
      ('equipment','eq','equipment',8,{'key':'key','owner':'owner','lat':'lat','lon':'lon','elev_ft':'elev_ft'}),
      ('anchors','anc','anchor',4,{'key':'key','owner':'owner','lat':'lat','lon':'lon','elev_ft':'elev_ft'}),
      ('crossarms','arm','crossarm',8,{'key':'key','owner':'owner','lat':'tip_lat','lon':'tip_lon','elev_ft':'elev_ft'}),
    ]
    for sheet,prefix,ftype,count,fields in specs:
        for r in sheets.get(sheet,[]):
            pid=as_int(r.get('internal_id')); pnum=val(r.get('pole_number'))
            for i in range(1,count+1):
                key=val(r.get(f'{prefix}_{i}_{fields.get("key","key")}')) if 'key' in fields else None
                owner=val(r.get(f'{prefix}_{i}_{fields.get("owner","owner")}')) if 'owner' in fields else None
                usage=val(r.get(f'{prefix}_{i}_{fields.get("usage","usage")}')) if 'usage' in fields else None
                lat=num(r.get(f'{prefix}_{i}_{fields["lat"]}')); lon=num(r.get(f'{prefix}_{i}_{fields["lon"]}')); elev=num(r.get(f'{prefix}_{i}_{fields["elev_ft"]}'))
                present=key or owner or usage or lat is not None or lon is not None or elev is not None
                if not present: continue
                if lat is not None and lon is not None and elev is not None:
                    x,y=tf.transform(lon,lat)
                    features.append(point_feature(pid,ftype,[x,y,elev],category=prefix,index=i,key=key,owner=owner,usage=usage,pole_number=pnum))
                elif ftype in ('equipment','anchor','crossarm'):
                    # no auto-fail for crossarms because SOW has conflicting instructions; review missing geometry only when partially populated
                    if any(v is not None for v in (lat,lon,elev)):
                        addfinding('PLA-R020','REVIEW',pid,pnum,sheet,f'{prefix}_{i}_geometry',f'{ftype.title()} {i} has partial XYZ geometry.',f'{lat},{lon},{elev}','complete XYZ or accepted blank')
                if ftype=='equipment' and key and not owner:
                    addfinding('PLA-R021','REVIEW',pid,pnum,sheet,f'{prefix}_{i}_owner',f'Equipment {i} has a key but owner is blank.',None,'owner or explicit UNKNOWN')

    # guys -> line to anchors when possible
    anchor_pts=defaultdict(dict)
    for f in features:
        if f['properties'].get('feature_type')=='anchor': anchor_pts[f['properties']['internal_id']][f['properties']['index']]=f['geometry']['coordinates']
    pole_lines={f['properties']['internal_id']:f for f in features if f['properties'].get('feature_type')=='pole'}
    for r in sheets['guys']:
        pid=as_int(r.get('internal_id')); pnum=val(r.get('pole_number'))
        for i in range(1,9):
            key=val(r.get(f'guy_{i}_key')); owner=val(r.get(f'guy_{i}_owner')); anc=as_int(r.get(f'guy_{i}_anchor_no')); elev=num(r.get(f'guy_{i}_elev_ft'))
            if not any(x is not None for x in (key,owner,anc,elev)): continue
            if anc and anc not in anchor_pts.get(pid,{}):
                addfinding('PLA-R010','FAIL',pid,pnum,'guys',f'guy_{i}_anchor_no',f'Guy {i} references missing anchor {anc}.',anc,'existing anchor number')
            if anc in anchor_pts.get(pid,{}) and elev is not None and pid in pole_lines:
                top=pole_lines[pid]['geometry']['coordinates'][1]
                start=[top[0],top[1],elev]
                features.append(line_feature(pid,'guy',start,anchor_pts[pid][anc],index=i,anchor_no=anc,key=key,owner=owner,pole_number=pnum))

    # sidewalk braces: missing guy_nos when brace exists
    for r in sheets['sidewalk_braces']:
        pid=as_int(r.get('internal_id')); pnum=val(r.get('pole_number'))
        for i in (1,2):
            key=val(r.get(f'swb_{i}_key')); guys=val(r.get(f'swb_{i}_guy_nos'))
            if key and not guys:
                addfinding('PLA-R012','FAIL',pid,pnum,'sidewalk_braces',f'swb_{i}_guy_nos',f'Sidewalk brace {i} exists but guy numbers are blank.',None,'one or more guy numbers')

    # span guy references: ensure runs_to token exists; explicit OTHER_n must be defined on poles row
    pole_rows={as_int(r.get('internal_id')):r for r in sheets['poles']}
    for r in sheets['span_guys']:
        pid=as_int(r.get('internal_id')); pnum=val(r.get('pole_number'))
        for i in range(1,6):
            key=val(r.get(f'sgy_{i}_key')); runs=val(r.get(f'sgy_{i}_runs_to')); end=val(r.get(f'sgy_{i}_end_type'))
            if not any((key,runs,end)): continue
            if runs and str(runs).upper().startswith('OTHER_'):
                n=as_int(str(runs).split('_')[-1]); prow=pole_rows.get(pid,{})
                if n and not (as_int(prow.get(f'other_{n}_id')) or val(prow.get(f'other_{n}_type'))):
                    addfinding('PLA-R030','FAIL',pid,pnum,'span_guys',f'sgy_{i}_runs_to',f'Span guy {i} references {runs}, but that Other connection is not defined.',runs,'defined Other connection')

    rules=[rule_status(rule_id,applicable) for rule_id,applicable in applies.items()]
    return {'poles':poles,'features':features,'findings':findings,'sheets':list(wb.sheetnames),'pole_sheet':pole_sheet.title,'rules':rules}


REQUIRED_POLE_COLUMNS=(('internal_id','Internal ID'),('pole_number','Pole Number'))


def _missing_columns_message(wb):
    """Which required columns the pole catalogue lacks, named the way the customer would look for them."""
    best=None
    for order,ws in enumerate(wb.worksheets):
        headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1),())]
        try: columns=header_map(headers)
        except ValueError as exc: return f"Sheet {ws.title!r}: {exc}. Keep only one such column."
        named=normalize_header(ws.title) in {'poles','pole','pole catalogue','pole catalog','zonename','zone name'}
        score=(named,len(columns),-order)
        if best is None or score>best[0]: best=(score,ws,columns)
    if best is None: return 'The workbook has no worksheets.'
    _,ws,columns=best
    missing=[label for key,label in REQUIRED_POLE_COLUMNS if key not in columns]
    if not has_location(columns): missing.append('Latitude and Longitude (or X and Y)')
    if not missing: return None
    return f"Sheet {ws.title!r} is missing required column{'s' if len(missing)>1 else ''}: {', '.join(missing)}."


def check_production_workbook(source):
    """Read only the pole catalogue of an Excel (path or bytes) and report what would stop or affect the import.

    errors   - the import cannot work (missing columns, rows without a Pole Number, bad or duplicate Internal IDs)
    warnings - it will import, but something needs attention (blank Internal IDs to set later, duplicate Pole Numbers,
               coordinates that are not numbers)"""
    from io import BytesIO
    result={'errors':[],'warnings':[],'pole_count':0,'missing_internal_ids':0,'sheet':None}
    try:
        wb=load_workbook(BytesIO(source) if isinstance(source,(bytes,bytearray)) else source,read_only=True,data_only=True)
    except Exception:
        result['errors'].append('This file could not be read as an Excel workbook (.xlsx).'); return result
    try:
        try:
            sheet,columns=detect_pole_sheet(wb,require_location=True)
        except ValueError as exc:
            result['errors'].append(_missing_columns_message(wb) or str(exc)); return result
        result['sheet']=sheet.title
        headers=[val(c.value) for c in next(sheet.iter_rows(min_row=1,max_row=1),())]
        label={key:str(headers[index]) for key,index in columns.items()}
        get=lambda row,key: val(row[columns[key]]) if key in columns and columns[key]<len(row) else None
        seen_ids={}; seen_numbers={}; blank_ids=[]; bad_coordinates=0
        for row_number,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
            if not any(val(v) is not None for v in row): continue
            raw_id,number=get(row,'internal_id'),get(row,'pole_number')
            if raw_id is None and number is None: continue
            result['pole_count']+=1
            if number is None:
                result['errors'].append(f"Row {row_number}: {label['pole_number']} is empty (Internal ID {raw_id}).")
            if raw_id is None:
                blank_ids.append(str(number))
            elif as_int(raw_id) is None:
                result['errors'].append(f"Row {row_number}: {label['internal_id']} {raw_id!r} is not a whole number.")
            elif as_int(raw_id) in seen_ids:
                result['errors'].append(f"Row {row_number}: {label['internal_id']} {as_int(raw_id)} is also used on row {seen_ids[as_int(raw_id)]}.")
            else:
                seen_ids[as_int(raw_id)]=row_number
            if number is not None:
                seen_numbers.setdefault(str(number),[]).append(row_number)
            for key in ('latitude','longitude','x','y'):
                raw=get(row,key)
                if raw is not None and num(raw) is None: bad_coordinates+=1
        if not result['pole_count']:
            result['errors'].append(f"Sheet {sheet.title!r} has the required columns but no pole rows.")
        if blank_ids:
            result['missing_internal_ids']=len(blank_ids)
            shown=', '.join(blank_ids[:8])+(' …' if len(blank_ids)>8 else '')
            result['warnings'].append(f"{len(blank_ids)} pole{'s have' if len(blank_ids)>1 else ' has'} no Internal ID ({shown}). "
                                      "They will be loaded; set each Internal ID in the Production view and it is saved to the Excel.")
        duplicates=[n for n,rows in seen_numbers.items() if len(rows)>1]
        if duplicates:
            result['warnings'].append(f"Pole Number used on more than one row: {', '.join(duplicates[:8])}{' …' if len(duplicates)>8 else ''}.")
        if bad_coordinates:
            result['warnings'].append(f"{bad_coordinates} coordinate cell{'s are' if bad_coordinates>1 else ' is'} not a number; those poles are reported in QC and are not placed on the LiDAR.")
        if len(result['errors'])>20:
            more=len(result['errors'])-20; result['errors']=result['errors'][:20]+[f"… and {more} more."]
        return result
    finally:
        wb.close()

