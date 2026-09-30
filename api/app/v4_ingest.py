"""v4 source-preserving workbook ingestion primitives."""
from __future__ import annotations
import hashlib
import json
from openpyxl import load_workbook

def stable_source_record_id(project_id: str, snapshot_id: str, sheet: str, row_number: int) -> str:
    raw = f"{project_id}|{snapshot_id}|{sheet}|{row_number}".encode("utf-8")
    return "SRC-" + hashlib.sha256(raw).hexdigest()[:24].upper()

def normalize_cell(value):
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value

def read_sheet_with_source_rows(path: str, sheet_name: str = "poles") -> list[dict]:
    """Return every non-empty source row, including rows with blank internal_id."""
    wb = load_workbook(path, read_only=True, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Workbook must contain {sheet_name!r} sheet")
    ws = wb[sheet_name]
    header = [normalize_cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {header[i]: normalize_cell(values[i]) for i in range(min(len(header), len(values))) if header[i]}
        if not any(v is not None for v in payload.values()):
            continue
        rows.append({"row_number": row_number, "payload": payload})
    return rows

def build_source_records(project_id: str, snapshot_id: str, path: str, sheet_name: str = "poles") -> list[dict]:
    out = []
    for row in read_sheet_with_source_rows(path, sheet_name):
        p = row["payload"]
        out.append({
            "id": stable_source_record_id(project_id, snapshot_id, sheet_name, row["row_number"]),
            "project_id": project_id,
            "snapshot_id": snapshot_id,
            "sheet_name": sheet_name,
            "row_number": row["row_number"],
            "pole_number": p.get("pole_number"),
            "source_internal_id": p.get("internal_id"),
            "source_payload_json": json.dumps(p, default=str, separators=(",", ":")),
        })
    return out
