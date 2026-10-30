"""Fill a PLA-1 Excel (one ZoneName pole sheet) with the pole values measured on the LiDAR.

Each pole's row (matched by Pole Number, as the Workbook Data tab does) gets, from its latest saved points:
  Pole_Base -> Latitude, Longitude, Bottom of Pole Elevation, First Pole, Last Pole, and the Incomplete Bottom /
               Latitude / Longitude of the spot toward an uncaptured pole (first_pole or last_pole INCOMPLETE)
  Pole_Top  -> Top of Pole Elevation
Other cells (Workbook Data edits, attachment columns) are left as they are. A pole whose Pole Number is on more than
one row is skipped and reported, because there is no telling which row is the pole. The source Excel is never
modified: the result is a new file built from the latest workbook bytes.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from .collection_workbook import restore_validation_extensions
from .source_schema import canonical_field, detect_pole_sheet, normalize_header
from .workbook_editor import normalize_pole_number

# Field -> accepted PLA-1 headers (normalized: lower case, punctuation and underscores as spaces).
POLE_COLUMNS = {
    "latitude": {"latitude"}, "longitude": {"longitude"},
    "bottom_elev": {"bottom of pole elevation"}, "top_elev": {"top of pole elevation"},
    "first_pole": {"first pole"}, "last_pole": {"last pole"},
    "inc_elev": {"incomplete bottom"}, "inc_lat": {"incomplete latitude"}, "inc_lon": {"incomplete longitude"},
}
# The point form's choices -> the PLA-1 dropdown ("YES, Incomplete").
CHOICES = {"YES": "YES", "INCOMPLETE": "Incomplete"}
POLE_FAMILIES = ("poles", "Pole Points")
BASE_FEATURES = ("Pole_Base", "Pole Base / Ground Point")
TOP_FEATURES = ("Pole_Top", "Pole Top Point")
# Older releases saved the incoming line as its own point instead of on the Pole_Base.
INCOMING_POINTS = {"prev_inc": "first_pole", "next_inc": "last_pole"}


def _number(value, digits):
    if value in (None, ""):
        return None
    try:
        return round(float(value), digits)
    except (TypeError, ValueError):
        return None


def pole_values(points: list[dict], elevation_to_ft: float = 1.0) -> dict:
    """{pole_internal_id: {field: value}} from each pole's latest Pole_Base / Pole_Top (and legacy incoming) points."""
    latest: dict[tuple, dict] = {}
    for point in sorted(points, key=lambda p: str(p.get("updated_at") or "")):
        if point.get("family") in POLE_FAMILIES and point.get("pole_internal_id") is not None:
            latest[(point["pole_internal_id"], point["feature_type"])] = point
    elevation = lambda z: _number(float(z) * elevation_to_ft, 3) if z not in (None, "") else None
    out: dict = {}
    for (pole_id, feature), point in latest.items():
        values = out.setdefault(pole_id, {})
        attrs = point.get("attributes") or {}
        if feature in BASE_FEATURES:
            values.update(latitude=_number(point.get("latitude"), 8), longitude=_number(point.get("longitude"), 8), bottom_elev=elevation(point.get("z")))
            for flag in ("first_pole", "last_pole"):
                choice = str(attrs.get(flag) or "").strip().upper()
                if choice:
                    values[flag] = CHOICES.get(choice, attrs[flag])
            # PLA-1 has one set of Incomplete columns: the spot toward the previous pole (first_pole INCOMPLETE), else the next.
            side = next((s for s, flag in INCOMING_POINTS.items() if str(attrs.get(flag) or "").upper() == "INCOMPLETE"), None)
            if side:
                for field, part, digits in (("inc_lat", "lat", 8), ("inc_lon", "lon", 8), ("inc_elev", "elev_ft", 3)):
                    number = _number(attrs.get(f"{side}_{part}"), digits)
                    if number is not None:
                        values[field] = number
        elif feature in TOP_FEATURES:
            values["top_elev"] = elevation(point.get("z"))
        elif feature in INCOMING_POINTS:
            values.setdefault("inc_lat", _number(point.get("latitude"), 8))
            values.setdefault("inc_lon", _number(point.get("longitude"), 8))
            values.setdefault("inc_elev", elevation(point.get("z")))
            values.setdefault(INCOMING_POINTS[feature], "Incomplete")
    return {pole_id: {k: v for k, v in values.items() if v is not None} for pole_id, values in out.items()}


def fill_pla1_workbook(base: bytes, pole_numbers: dict, points: list[dict], *, elevation_to_ft: float = 1.0) -> tuple[bytes, dict]:
    """Return the PLA-1 workbook bytes with every pole's LiDAR values written in, and a summary."""
    summary = {"poles_filled": 0, "cells_written": 0, "skipped": [], "missing_columns": []}
    wb = load_workbook(BytesIO(base))
    sheet, _ = detect_pole_sheet(wb, require_location=False)
    headers = [cell.value for cell in sheet[1]]
    columns: dict[str, int] = {}
    pole_column = None
    for index, header in enumerate(headers, 1):
        name = normalize_header(header)
        if pole_column is None and canonical_field(header) == "pole_number":
            pole_column = index
        for field, names in POLE_COLUMNS.items():
            if name in names and field not in columns:
                columns[field] = index
    rows: dict[str, list[int]] = {}
    if pole_column:
        for row in range(2, sheet.max_row + 1):
            key = normalize_pole_number(sheet.cell(row, pole_column).value)
            if key:
                rows.setdefault(key, []).append(row)
    missing = set()
    for pole_id, values in pole_values(points, elevation_to_ft).items():
        if not values:
            continue
        number = pole_numbers.get(pole_id)
        matches = rows.get(normalize_pole_number(number)) or []
        if not matches:
            summary["skipped"].append(f"{number or f'Internal ID {pole_id}'}: not in the Excel")
            continue
        if len(matches) > 1:
            summary["skipped"].append(f"{number}: on {len(matches)} rows of the Excel")
            continue
        written = 0
        for field, value in values.items():
            if field not in columns:
                missing.add(field)
                continue
            sheet.cell(matches[0], columns[field]).value = value
            written += 1
        if written:
            summary["poles_filled"] += 1
            summary["cells_written"] += written
    summary["missing_columns"] = sorted(missing)
    out = BytesIO()
    wb.save(out)
    return restore_validation_extensions(base, out.getvalue()), summary
