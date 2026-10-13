"""Write saved production points into the COLLECTION workbook layout.

One sheet per group, one row per pole (keyed by ``internal_id``), and one numbered column set per point, e.g.
``swb_1_key … swb_1_guy_nos, swb_2_key …``. A point whose number has no columns yet gets a new set inserted after
the last set of its prefix, with the same field names, header style, widths, table columns and dropdowns.

The source workbook is never modified: the result is a new file built from the latest workbook bytes plus the points.
openpyxl drops Excel's x14 "extension" data validations (lookup-list dropdowns) on save, so they are carried over
from the original XML afterwards, shifted for inserted columns and copied onto new column sets.
"""
from __future__ import annotations

import re
import zipfile
from copy import copy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.table import TableColumn

# Column template of the COLLECTION workbook (names only): fixed columns, numbered groups, trailing columns.
LAYOUT = {
    "poles": {
        "before": ["internal_id", "pole_number", "fqn_id", "attach_identifying_no", "attach_location_on_pole", "attach_status",
                   "abandoned", "attach_description", "voltage", "pole_owner_raw", "grade_of_construction", "first_pole",
                   "prev_inc_lat", "prev_inc_lon", "prev_inc_elev_ft", "last_pole", "next_inc_lat", "next_inc_lon",
                   "next_inc_elev_ft", "latitude", "longitude", "bottom_elev_ft", "top_lat", "top_lon", "top_elev_ft", "environment"],
        "groups": {"other": (["type", "id", "attach_N", "lat", "lon", "elev_ft"], 4)},
    },
    "attachments": {"groups": {"comm": (["owner", "size", "support", "arm_no", "lat", "lon", "elev_ft"], 9),
                               "util": (["usage", "owner", "size", "support", "arm_no", "lat", "lon", "elev_ft"], 21)}},
    "crossarms": {"groups": {"arm": (["key", "owner", "tip_lat", "tip_lon", "elev_ft"], 8)}},
    "equipment": {"groups": {"eq": (["key", "owner", "lat", "lon", "elev_ft"], 8)}},
    "anchors": {"groups": {"anc": (["key", "owner", "lat", "lon", "elev_ft"], 4)}},
    "guys": {"groups": {"guy": (["key", "owner", "anchor_no", "elev_ft"], 8)}},
    "sidewalk_braces": {"groups": {"swb": (["key", "owner", "anchor_no", "elev_ft", "length_ft", "guy_nos"], 2)}},
    "span_guys": {"groups": {"sgy": (["key", "owner", "runs_to", "end_type", "elev_ft"], 3)}},
}
# Production annotation group → (sheet, column prefix). Point numbers come from the feature type (arm_3 → 3).
FAMILY_SHEETS = {
    "crossarms": ("crossarms", "arm"), "attachments_comm": ("attachments", "comm"), "attachments_util": ("attachments", "util"),
    "equipment": ("equipment", "eq"), "anchors": ("anchors", "anc"), "guys": ("guys", "guy"),
    "span_guys": ("span_guys", "sgy"), "sidewalk_braces": ("sidewalk_braces", "swb"), "other_poles": ("poles", "other"),
}
POLE_BASE_FIELDS = ("first_pole", "last_pole", "environment")
OTHER_ON_POLE_BASE = {"other_type": "type", "other_id": "id", "other_attach_N": "attach_N"}
HEADER_FILL = PatternFill("solid", fgColor="FF375623")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
X14_URI = "{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}"


def _key(value):
    if value is None:
        return None
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def is_collection_workbook(wb) -> bool:
    if "poles" not in wb.sheetnames or not any(sheet in wb.sheetnames for sheet, _ in FAMILY_SHEETS.values() if sheet != "poles"):
        return False
    return "internal_id" in [c.value for c in wb["poles"][1]]


def _shift_ref(ref: str, pos: int, amount: int) -> str:
    """Shift every column >= pos in a space-separated range list ("C2:C1111 H2:H1111")."""
    out = []
    for part in ref.split():
        cells = []
        for cell in part.split(":"):
            m = re.match(r"^(\$?)([A-Z]+)(\$?)(\d*)$", cell)
            if not m:
                cells.append(cell)
                continue
            col = column_index_from_string(m.group(2))
            cells.append(f"{m.group(1)}{get_column_letter(col + amount if col >= pos else col)}{m.group(3)}{m.group(4)}")
        out.append(":".join(cells))
    return " ".join(out)


def _copy_ref(ref: str, template_col: int, new_col: int) -> str:
    """Add, for each range that covers the template column alone, the same rows on the new column."""
    added = []
    for part in ref.split():
        cells = part.split(":")
        cols = []
        for cell in cells:
            m = re.match(r"^\$?([A-Z]+)\$?(\d*)$", cell)
            cols.append((column_index_from_string(m.group(1)), m.group(2)) if m else (None, ""))
        if all(c == template_col for c, _ in cols):
            added.append(":".join(f"{get_column_letter(new_col)}{rows}" for _, rows in cols))
    return " ".join(ref.split() + added) if added else ref


class _Sheet:
    def __init__(self, ws, ops: list):
        self.ws, self.ops = ws, ops
        self.headers = [c.value for c in ws[1]]
        self.col = {str(h): i + 1 for i, h in enumerate(self.headers) if h is not None}
        id_col = self.col.get("internal_id", 1)
        self.rows = {}
        for r in range(2, ws.max_row + 1):
            k = _key(ws.cell(r, id_col).value)
            if k is not None and k not in self.rows:
                self.rows[k] = r

    def set_numbers(self, prefix: str) -> list[int]:
        return sorted({int(m.group(1)) for h in self.col if (m := re.match(rf"^{prefix}_(\d+)_", h))})

    def suffixes(self, prefix: str) -> list[str]:
        numbers = self.set_numbers(prefix)
        if numbers:
            first = numbers[0]
            return [h[len(f"{prefix}_{first}_"):] for h in sorted(self.col, key=self.col.get) if h.startswith(f"{prefix}_{first}_")]
        for spec in LAYOUT.get(self.ws.title, {}).get("groups", {}).items():
            if spec[0] == prefix:
                return list(spec[1][0])
        return ["key", "owner", "lat", "lon", "elev_ft"]

    def ensure_set(self, prefix: str, n: int, summary: dict) -> None:
        suffixes = self.suffixes(prefix)
        if f"{prefix}_{n}_{suffixes[0]}" in self.col:
            return
        numbers = self.set_numbers(prefix)
        start = (max(numbers) + 1) if numbers else 1
        for m in range(start, n + 1):
            self._insert_set(prefix, m, suffixes, numbers[-1] if numbers else None)
            summary["columns_added"].append(f"{self.ws.title}: {prefix}_{m}_* ({len(suffixes)} columns)")
            numbers = self.set_numbers(prefix)

    def _insert_set(self, prefix: str, m: int, suffixes: list[str], template_n: int | None) -> None:
        ws = self.ws
        group_cols = [c for h, c in self.col.items() if re.match(rf"^{prefix}_\d+_", h)]
        if group_cols:
            pos = max(group_cols) + 1
        elif "remarks" in self.col:
            pos = self.col["remarks"]
        else:
            pos = ws.max_column + 1
        amount = len(suffixes)
        widths = {h: ws.column_dimensions[get_column_letter(c)].width for h, c in self.col.items()}
        template_cols = [self.col.get(f"{prefix}_{template_n}_{s}") if template_n else None for s in suffixes]
        if pos <= ws.max_column:
            ws.insert_cols(pos, amount)
        self.ops.append(("insert", pos, amount))
        # Bookkeeping for columns that moved.
        self.col = {h: (c + amount if c >= pos else c) for h, c in self.col.items()}
        for i, suffix in enumerate(suffixes):
            new_col, name = pos + i, f"{prefix}_{m}_{suffix}"
            cell = ws.cell(1, new_col, name)
            template = template_cols[i]
            if template:
                source = ws.cell(1, template)
                cell.font, cell.fill, cell.border, cell.alignment = copy(source.font), copy(source.fill), copy(source.border), copy(source.alignment)
                cell.number_format = source.number_format
            else:
                cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            self.col[name] = new_col
            if template:
                self.ops.append(("copy", template, new_col))
        for h, c in self.col.items():
            width = widths.get(h) or (widths.get(f"{prefix}_{template_n}_{h.split('_', 2)[-1]}") if template_n else None)
            if width:
                ws.column_dimensions[get_column_letter(c)].width = width
        self.headers = [c.value for c in ws[1]]
        self._grow_tables(pos, amount, suffixes, m, prefix)
        self._shift_validations(pos, amount, template_cols)

    def _grow_tables(self, pos: int, amount: int, suffixes: list[str], m: int, prefix: str) -> None:
        for table in self.ws.tables.values():
            first, last = table.ref.split(":")
            c1 = column_index_from_string(re.match(r"[A-Z]+", first).group(0))
            c2 = column_index_from_string(re.match(r"[A-Z]+", last).group(0))
            if not (c1 <= pos <= c2 + 1):
                continue
            last_row = re.search(r"\d+$", last).group(0)
            table.ref = f"{first}:{get_column_letter(c2 + amount)}{last_row}"
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
            columns = list(table.tableColumns)
            index = pos - c1
            for i, suffix in enumerate(suffixes):
                columns.insert(index + i, TableColumn(id=0, name=f"{prefix}_{m}_{suffix}"))
            for i, column in enumerate(columns, start=1):
                column.id = i
            table.tableColumns = columns

    def _shift_validations(self, pos: int, amount: int, template_cols: list) -> None:
        for dv in self.ws.data_validations.dataValidation:
            ref = _shift_ref(str(dv.sqref), pos, amount)
            for i, template in enumerate(template_cols):
                if template:
                    ref = _copy_ref(ref, template, pos + i)
            dv.sqref = ref

    def add_pole_rows(self, poles: list[dict]) -> int:
        """Rows for poles the sheet does not list yet, at the bottom (inside its table)."""
        ws, added = self.ws, 0
        formula = None
        if "pole_number" in self.col:
            sample = ws.cell(2, self.col["pole_number"]).value if ws.max_row >= 2 else None
            formula = sample if isinstance(sample, str) and sample.startswith("=") else None
        # Template sheets come with pre-formatted empty rows (internal_id blank, pole_number formula waiting):
        # fill those first, in order, and only append below when they run out.
        id_col = self.col.get("internal_id", 1)
        empty = [r for r in range(2, ws.max_row + 1) if ws.cell(r, id_col).value in (None, "")]
        for pole in sorted(poles, key=lambda p: (str(type(p["internal_id"])), p["internal_id"])):
            key = _key(pole["internal_id"])
            if key in self.rows:
                continue
            r = empty.pop(0) if empty else max(ws.max_row, 1) + 1
            if r <= ws.max_row and ws.cell(r, self.col.get("pole_number", 2)).value not in (None, ""):
                ws.cell(r, id_col, pole["internal_id"])
                self.rows[key] = r
                added += 1
                continue
            ws.cell(r, self.col.get("internal_id", 1), pole["internal_id"])
            if "pole_number" in self.col:
                # Keep the sheet's own pole_number lookup formula when it has one; otherwise write the number.
                ws.cell(r, self.col["pole_number"], formula.replace("$A2", f"$A{r}") if formula and "$A2" in formula else pole.get("pole_number"))
            self.rows[key] = r
            added += 1
        if added:
            for table in ws.tables.values():
                first, last = table.ref.split(":")
                table.ref = f"{first}:{re.match(r'[A-Z]+', last).group(0)}{ws.max_row}"
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
        return added

    def write(self, row: int, column: str, value, summary: dict) -> bool:
        col = self.col.get(column)
        if not col:
            summary["unmapped"].add(f"{self.ws.title}: {column}")
            return False
        self.ws.cell(row, col, value)
        return True


def _generated_workbook(poles: list[dict]):
    wb = Workbook()
    wb.remove(wb.active)
    for name, spec in LAYOUT.items():
        ws = wb.create_sheet(name)
        headers = list(spec.get("before", ["internal_id", "pole_number"]))
        for prefix, (suffixes, slots) in spec["groups"].items():
            headers += [f"{prefix}_{n}_{s}" for n in range(1, slots + 1) for s in suffixes]
        headers.append("remarks")
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(1, c, h)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            ws.column_dimensions[get_column_letter(c)].width = max(12, min(28, len(h) + 2))
        for r, pole in enumerate(sorted(poles, key=lambda p: p["internal_id"]), start=2):
            ws.cell(r, 1, pole["internal_id"])
            ws.cell(r, 2, pole.get("pole_number"))
        ws.freeze_panes = "C2"
    return wb


def build_collection_workbook(base: bytes | None, poles: list[dict], points: list[dict], *, elevation_to_ft: float = 1.0) -> tuple[bytes, dict]:
    """Return the workbook bytes with every point written in, and a summary of what was done."""
    summary = {"source": "generated", "points_written": 0, "columns_added": [], "rows_added": 0, "skipped": [], "unmapped": set()}
    wb = None
    if base:
        try:
            wb = load_workbook(BytesIO(base))
        except Exception:  # not a readable .xlsx; fall back to a generated layout
            wb = None
        if wb is not None and not is_collection_workbook(wb):
            wb = None
        elif wb is not None:
            summary["source"] = "workbook"
    if wb is None:
        wb = _generated_workbook(poles)
        base = None
    ops: dict[str, list] = {}
    sheets: dict[str, _Sheet] = {}
    # Every pole gets a row on every group sheet, whether or not it has points yet.
    for name in LAYOUT:
        if name in wb.sheetnames:
            sheets[name] = _Sheet(wb[name], ops.setdefault(name, []))
            summary["rows_added"] += sheets[name].add_pole_rows(poles)

    def sheet(name: str) -> _Sheet | None:
        if name not in wb.sheetnames:
            return None
        if name not in sheets:
            sheets[name] = _Sheet(wb[name], ops.setdefault(name, []))
            summary["rows_added"] += sheets[name].add_pole_rows(poles)
        return sheets[name]

    elevation = lambda z: round(float(z) * elevation_to_ft, 3) if z is not None else None
    latest: dict[tuple, dict] = {}
    for point in sorted(points, key=lambda p: str(p.get("updated_at") or "")):
        latest[(point.get("pole_internal_id"), point["family"], point["feature_type"])] = point

    for (pole_id, family, feature), point in latest.items():
        if pole_id is None:
            summary["skipped"].append(f"{feature}: no pole")
            continue
        attrs = point.get("attributes") or {}
        lat, lon, z = point.get("latitude"), point.get("longitude"), point.get("z")
        lat = round(float(lat), 8) if lat is not None else None
        lon = round(float(lon), 8) if lon is not None else None
        if family in ("poles", "Pole Points"):
            s = sheet("poles")
            if not s:
                continue
            row = s.rows.get(_key(pole_id))
            is_base = feature in ("Pole_Base", "Pole Base / Ground Point")
            is_top = feature in ("Pole_Top", "Pole Top Point")
            if is_base:
                s.write(row, "latitude", lat, summary); s.write(row, "longitude", lon, summary); s.write(row, "bottom_elev_ft", elevation(z), summary)
                for field in POLE_BASE_FIELDS:
                    if attrs.get(field) not in (None, ""):
                        s.write(row, field, attrs[field], summary)
                if any(attrs.get(k) not in (None, "") for k in OTHER_ON_POLE_BASE):
                    s.ensure_set("other", 1, summary)
                    for key, suffix in OTHER_ON_POLE_BASE.items():
                        if attrs.get(key) not in (None, ""):
                            s.write(row, f"other_1_{suffix}", attrs[key], summary)
            elif is_top:
                s.write(row, "top_lat", lat, summary); s.write(row, "top_lon", lon, summary); s.write(row, "top_elev_ft", elevation(z), summary)
            else:
                summary["skipped"].append(f"{feature}: not a workbook column")
                continue
            summary["points_written"] += 1
            continue
        if family not in FAMILY_SHEETS:
            summary["skipped"].append(f"{feature}: group {family} has no workbook sheet")
            continue
        sheet_name, prefix = FAMILY_SHEETS[family]
        m = re.match(rf"^{prefix}_(\d+)$", str(feature))
        s = sheet(sheet_name)
        if not m or not s:
            summary["skipped"].append(f"{feature}: no numbered column set")
            continue
        n = int(m.group(1))
        s.ensure_set(prefix, n, summary)
        row = s.rows.get(_key(pole_id))
        for suffix in s.suffixes(prefix):
            value = {"lat": lat, "tip_lat": lat, "lon": lon, "tip_lon": lon, "elev_ft": elevation(z)}.get(suffix, ...)
            if value is not ...:
                s.write(row, f"{prefix}_{n}_{suffix}", value, summary)
        for key, value in attrs.items():
            if not key.startswith(f"{prefix}_") or value in (None, ""):
                continue
            field = re.sub(r"^\d+_", "", key[len(prefix) + 1:])
            s.write(row, f"{prefix}_{n}_{field}", value, summary)
        summary["points_written"] += 1

    if base:
        wb.calculation.fullCalcOnLoad = True  # pole_number lookups recalculate when opened
    out = BytesIO()
    wb.save(out)
    data = out.getvalue()
    if base:
        data = restore_validation_extensions(base, data, ops)
    summary["unmapped"] = sorted(summary["unmapped"])
    return data, summary


def _sheet_paths(archive: zipfile.ZipFile) -> dict[str, str]:
    wb = archive.read("xl/workbook.xml").decode("utf-8")
    rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for rel in re.findall(r"<Relationship\b[^>]*/>", rels):
        rid, target = re.search(r'Id="([^"]+)"', rel), re.search(r'Target="([^"]+)"', rel)
        if rid and target:
            t = target.group(1).lstrip("/")
            targets[rid.group(1)] = t if t.startswith("xl/") else f"xl/{t}"
    paths = {}
    for tag in re.findall(r"<sheet\b[^>]*/>", wb):
        name, rid = re.search(r'name="([^"]+)"', tag), re.search(r'r:id="([^"]+)"', tag)
        if name and rid and rid.group(1) in targets:
            paths[name.group(1)] = targets[rid.group(1)]
    return paths


def restore_validation_extensions(original: bytes, updated: bytes, ops: dict[str, list] | None = None) -> bytes:
    """Carry x14 data validations (dropdowns that point at other sheets) from the original into the saved file."""
    ops = ops or {}
    try:
        src, dst = zipfile.ZipFile(BytesIO(original)), zipfile.ZipFile(BytesIO(updated))
        src_paths, dst_paths = _sheet_paths(src), _sheet_paths(dst)
    except (KeyError, zipfile.BadZipFile):
        return updated
    replacements = {}
    for name, src_path in src_paths.items():
        if name not in dst_paths:
            continue
        xml = src.read(src_path).decode("utf-8")
        block = re.search(rf'<ext\b[^>]*uri="{re.escape(X14_URI)}"[^>]*>.*?</ext>', xml, re.S)
        if not block:
            continue
        ext = re.sub(r'\s+xr\d*:[A-Za-z]+="[^"]*"', "", block.group(0))
        def fix(match):
            ref = match.group(1)
            for op in ops.get(name, []):
                ref = _shift_ref(ref, op[1], op[2]) if op[0] == "insert" else _copy_ref(ref, op[1], op[2])
            return f"<xm:sqref>{ref}</xm:sqref>"
        ext = re.sub(r"<xm:sqref>([^<]+)</xm:sqref>", fix, ext)
        target = dst.read(dst_paths[name]).decode("utf-8")
        if "<extLst>" in target:
            target = target.replace("<extLst>", f"<extLst>{ext}", 1)
        else:
            target = target.replace("</worksheet>", f"<extLst>{ext}</extLst></worksheet>")
        replacements[dst_paths[name]] = target.encode("utf-8")
    if not replacements:
        return updated
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in dst.infolist():
            z.writestr(item, replacements.get(item.filename, dst.read(item.filename)))
    return out.getvalue()
