"""One-click deliverable package for an approved dataset version.

The ZIP holds the updated Production Excel, the QC Excel that was checked, the saved Production points as GeoJSON,
an Excel findings report, a README and a manifest with SHA-256 checksums. It is built when the version is approved
and stored as a derived object, so later edits never change what was delivered. Source objects are never modified.
"""
from __future__ import annotations

import hashlib
import io
import json
import pathlib
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .models import AuditLog, DatasetFile, Finding, Pole, ProductionAnnotation, Project, ReviewDecision
from .pipeline import dataset_pair
from .production import annotations_feature_collection
from .storage import object_exists, read_bytes, upload_bytes
from .workflow import CorrectionRequest, DatasetVersion, FindingRevision, version_files

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="0D2D46")


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text or "").strip("._") or "dataset"


def package_keys(project: Project, version: DatasetVersion) -> tuple[str, str]:
    base = f"{project.id}/versions/v{version.version_no}/deliverables/{_slug(project.name)}-v{version.version_no}-deliverable"
    return f"{base}.zip", f"{base}.manifest.json"


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value else None


def _latest_version(db: Session, project_id: str) -> DatasetVersion | None:
    return db.query(DatasetVersion).filter_by(project_id=project_id).order_by(DatasetVersion.version_no.desc()).first()


def _workbook(db: Session, version: DatasetVersion | None, roles: tuple[str, ...], cutoff: datetime | None) -> DatasetFile | None:
    """Newest uploaded file of the first role found, ignoring files created after ``cutoff`` (the approval time)."""
    if not version:
        return None
    files = [f for f in version_files(db, version.id) if f.status == "UPLOADED"]
    for role in roles:
        candidates = [f for f in files if f.role == role and (cutoff is None or f.created_at is None or _naive(f.created_at) <= _naive(cutoff))]
        if candidates:
            return max(candidates, key=lambda f: (_naive(f.created_at) if f.created_at else datetime.min))
    if cutoff is not None:
        # Nothing predates the approval (e.g. clock skew): deliver the original import rather than no Excel at all.
        return _workbook(db, version, roles[-1:], None)
    return None


def _naive(value: datetime) -> datetime:
    return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value


def _approval(db: Session, version: DatasetVersion) -> dict:
    row = (db.query(AuditLog).filter(AuditLog.action == "APPROVE_DATASET_VERSION", AuditLog.entity_id == version.id)
           .order_by(AuditLog.created_at.desc()).first())
    return {"approved_by": row.actor if row else None, "approved_at": _iso(version.approved_at) or (_iso(row.created_at) if row else None)}


def _findings(db: Session, project: Project, version: DatasetVersion) -> tuple[list[dict], bool]:
    """Findings of the approved version with the latest reviewer decision; older versions use their stored snapshots."""
    latest = _latest_version(db, project.id)
    if latest and latest.id == version.id:
        rows = db.query(Finding).filter_by(project_id=project.id).order_by(Finding.internal_id, Finding.id).all()
        decisions: dict[str, ReviewDecision] = {}
        for decision in db.query(ReviewDecision).filter(ReviewDecision.finding_id.in_([f.id for f in rows])).order_by(ReviewDecision.created_at).all() if rows else []:
            decisions[decision.finding_id] = decision
        out = []
        for f in rows:
            d = decisions.get(f.id)
            out.append({"id": f.id, "rule_id": f.rule_id, "severity": f.severity, "pole_number": f.pole_number, "internal_id": f.internal_id,
                        "sheet": f.sheet, "field": f.field, "message": f.message, "actual": f.actual, "expected": f.expected, "status": f.status,
                        "decision": d.decision if d else None, "reviewer": d.reviewer_email if d else None,
                        "decided_at": _iso(d.created_at) if d else None, "decision_comment": d.comment if d else None})
        return out, False
    snapshots = [json.loads(r.snapshot_json or "{}") for r in db.query(FindingRevision).filter_by(version_id=version.id).order_by(FindingRevision.id).all()]
    return [{**s, "decision": None, "reviewer": None, "decided_at": None, "decision_comment": None} for s in snapshots], True


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for values in rows:
        sheet.append([None if v is None else v for v in values])
    # Uploaded workbook text is untrusted: store every string as text so "=..." can never become a formula.
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1])) for r in rows[:500] if r[index - 1] is not None])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, width + 2), 60)
    return sheet


def findings_report(db: Session, project: Project, version: DatasetVersion, production: Project | None, approval: dict, generated_at: str) -> bytes:
    findings, historical = _findings(db, project, version)
    corrections = db.query(CorrectionRequest).filter_by(project_id=project.id, version_id=version.id).order_by(CorrectionRequest.created_at).all()
    by_id = {f["id"]: f for f in findings}
    poles = db.query(Pole).filter_by(project_id=project.id).order_by(Pole.internal_id).all()
    severity = Counter(f["severity"] for f in findings)
    status = Counter(f["status"] for f in findings)
    workbook = Workbook()
    summary = [
        ["Dataset", project.name], ["Dataset ID", project.id], ["Production dataset", production.name if production else "—"],
        ["Version", f"v{version.version_no}"], ["Version status", version.status],
        ["Approved by", approval["approved_by"] or "—"], ["Approved at (UTC)", approval["approved_at"] or "—"],
        ["Report generated at (UTC)", generated_at], ["Coordinate system", project.crs],
        ["Poles", len(poles)], ["Poles PASS", sum(p.qc_status == "PASS" for p in poles)], ["Poles FAIL", sum(p.qc_status == "FAIL" for p in poles)],
        ["Findings", len(findings)], ["FAIL", severity.get("FAIL", 0)], ["REVIEW", severity.get("REVIEW", 0)], ["UNVERIFIABLE", severity.get("UNVERIFIABLE", 0)],
        ["Findings open", status.get("OPEN", 0)], ["Corrections", len(corrections)], ["Corrections still open", sum(c.status == "OPEN" for c in corrections)],
    ]
    if historical:
        summary.append(["Note", "This version is no longer the latest; findings come from the snapshot stored with its QC run, without reviewer decisions."])
    _write_sheet(workbook, "Summary", ["Item", "Value"], summary, first=True)
    _write_sheet(workbook, "Findings",
                 ["Finding ID", "Rule", "Severity", "Pole Number", "internal_id", "Sheet", "Field", "Message", "Actual", "Expected", "Status",
                  "Reviewer decision", "Reviewer", "Decided at (UTC)", "Decision comment"],
                 [[f.get("id"), f.get("rule_id"), f.get("severity"), f.get("pole_number"), f.get("internal_id"), f.get("sheet"), f.get("field"),
                   f.get("message"), f.get("actual"), f.get("expected"), f.get("status"), f.get("decision"), f.get("reviewer"),
                   f.get("decided_at"), f.get("decision_comment")] for f in findings])
    _write_sheet(workbook, "Corrections",
                 ["Correction ID", "Finding ID", "Pole Number", "Rule", "Status", "Comment", "Requested by", "Requested at (UTC)", "Resolved by", "Resolved at (UTC)"],
                 [[c.id, c.finding_id, (by_id.get(c.finding_id) or {}).get("pole_number"), (by_id.get(c.finding_id) or {}).get("rule_id"), c.status,
                   c.comment, c.created_by, _iso(c.created_at), c.resolved_by, _iso(c.resolved_at)] for c in corrections])
    _write_sheet(workbook, "Poles", ["internal_id", "Pole Number", "QC status", "FAIL", "REVIEW", "UNVERIFIABLE"],
                 [[p.internal_id, p.pole_number, p.qc_status, p.qc_fail, p.qc_review, p.qc_unverifiable] for p in poles])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


README = """Deliverable package
===================

{name} — version v{version}
Approved by {approved_by} at {approved_at} (UTC). Package generated {generated_at} (UTC){late}.

Contents
{listing}

manifest.json lists every file with its SHA-256 checksum and the stored object it came from.
Coordinates in annotations.geojson are WGS84 longitude/latitude; Z is in the project units ({units}).
"""


def build_package(db: Session, project: Project, version: DatasetVersion, actor: str, at_approval: bool) -> dict:
    production, _ = dataset_pair(db, project)
    approval = _approval(db, version)
    cutoff = version.approved_at
    generated_at = datetime.now(timezone.utc).isoformat()
    entries: list[tuple[str, bytes, str, str | None]] = []  # (zip path, bytes, description, source object key)

    if production:
        source = _workbook(db, _latest_version(db, production.id), ("WORKBOOK_EDITED", "WORKBOOK"), cutoff)
        if source:
            label = "updated in Production" if source.role == "WORKBOOK_EDITED" else "as imported (no Production edits saved)"
            entries.append((f"excel/{_slug(pathlib.Path(source.filename).stem)}-production.xlsx", read_bytes(source.object_key),
                            f"Production Excel, {label}", source.object_key))
    checked = _workbook(db, version, ("WORKBOOK",), None)
    if checked:
        entries.append((f"excel/qc/{_slug(pathlib.Path(checked.filename).stem)}.xlsx", read_bytes(checked.object_key),
                        "Excel that QC checked for this version" if production else "Excel for this version", checked.object_key))
    if production:
        rows = db.query(ProductionAnnotation).filter_by(project_id=production.id).order_by(ProductionAnnotation.created_at, ProductionAnnotation.id).all()
        geo = annotations_feature_collection(db, production, rows)
        entries.append(("annotations/annotations.geojson", json.dumps(geo, ensure_ascii=False, indent=1).encode("utf-8"),
                        f"{len(rows)} saved Production points", None))
    entries.append(("reports/findings-report.xlsx", findings_report(db, project, version, production, approval, generated_at),
                    "Findings, reviewer decisions, corrections and pole QC status", None))

    files = [{"path": path, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), "description": text, "source_object_key": key}
             for path, data, text, key in entries]
    zip_key, manifest_key = package_keys(project, version)
    manifest = {
        "package": pathlib.Path(zip_key).name, "package_key": zip_key, "project_id": project.id, "project_name": project.name,
        "production_project_id": production.id if production else None, "version_id": version.id, "version_no": version.version_no,
        **approval, "generated_at": generated_at, "generated_by": actor, "built_at_approval": at_approval, "files": files,
    }
    listing = "\n".join(f"  {f['path']:<52} {f['description']}" for f in files)
    readme = README.format(name=project.name, version=version.version_no, approved_by=approval["approved_by"] or "—",
                           approved_at=approval["approved_at"] or "—", generated_at=generated_at, units=project.units,
                           late="" if at_approval else "; built after approval from the data at download time", listing=listing)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path, data, _, _ in entries:
            archive.writestr(path, data)
        archive.writestr("README.txt", readme)
        archive.writestr("manifest.json", json.dumps(manifest, indent=2))
    package = buffer.getvalue()
    manifest["package_sha256"] = hashlib.sha256(package).hexdigest()
    manifest["package_bytes"] = len(package)
    upload_bytes(package, zip_key, "application/zip")
    upload_bytes(json.dumps(manifest, indent=2).encode("utf-8"), manifest_key, "application/json")
    db.add(AuditLog(actor=actor, action="BUILD_DELIVERABLE_PACKAGE", entity_type="dataset_version", entity_id=version.id,
                    detail_json=json.dumps({"project_id": project.id, "package_key": zip_key, "sha256": manifest["package_sha256"],
                                            "bytes": len(package), "built_at_approval": at_approval})))
    return manifest


def ensure_package(db: Session, project: Project, version: DatasetVersion, actor: str) -> dict:
    """The stored package for this approved version, built now if it was approved before packages existed."""
    _, manifest_key = package_keys(project, version)
    if object_exists(manifest_key):
        return json.loads(read_bytes(manifest_key))
    return build_package(db, project, version, actor, at_approval=False)
