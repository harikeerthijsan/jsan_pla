from __future__ import annotations

import math
import re
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook
from .source_schema import canonical_field, detect_pole_sheet, normalize_header


def normalize_pole_number(value) -> str | None:
    if value is None or isinstance(value, (dict, list, bool)):
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer():
            value = int(value)
    text = str(value).strip().upper()
    if re.fullmatch(r"-?\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text or None


def _serializable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def match_geojson_pole_number(features: list[dict], pole_number, internal_id=None) -> dict:
    target = normalize_pole_number(pole_number)
    target_internal = normalize_pole_number(internal_id)
    if target is None and target_internal is None:
        return {"status": "MISSING", "feature_indexes": []}
    matching = []
    for index, feature in enumerate(features):
        if feature.get("geometry_type") != "Point":
            continue
        properties = feature.get("properties") or {}
        pole_values = [normalize_pole_number(value) for key, value in properties.items() if canonical_field(key) == "pole_number"]
        internal_values = [normalize_pole_number(value) for key, value in properties.items() if canonical_field(key) == "internal_id"]
        pole_values = [value for value in pole_values if value]
        internal_values = [value for value in internal_values if value]
        pole_match = target is not None and target in pole_values
        internal_match = target_internal is not None and target_internal in internal_values
        conflict = (pole_values and target is not None and not pole_match) or (internal_values and target_internal is not None and not internal_match)
        if not conflict and (pole_match or internal_match):
            matching.append(index)
    return {
        "status": "MISSING" if not matching else "MATCHED" if len(matching) == 1 else "AMBIGUOUS",
        "feature_indexes": matching,
    }


def _header_columns(worksheet) -> list[tuple[int, str]]:
    headers = []
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(1, column).value
        if value is not None and str(value).strip():
            headers.append((column, str(value)))
    return headers


def _matching_rows(worksheet, pole_number, internal_id=None) -> tuple[list[tuple[int, str]], list[int], int, list[int]]:
    headers = _header_columns(worksheet)
    pole_columns = [column for column, label in headers if canonical_field(label) == "pole_number"]
    internal_columns = [column for column, label in headers if canonical_field(label) == "internal_id"]
    if len(pole_columns) > 1:
        raise ValueError(f"Worksheet {worksheet.title!r} has multiple Pole Number columns")
    if len(internal_columns) > 1:
        raise ValueError(f"Worksheet {worksheet.title!r} has multiple Internal ID columns")
    if not pole_columns and not internal_columns:
        return headers, [], 0, []
    pole_column = pole_columns[0] if pole_columns else 0
    internal_column = internal_columns[0] if internal_columns else 0
    target = normalize_pole_number(pole_number)
    target_internal = normalize_pole_number(internal_id)
    rows = []
    for row in range(2, worksheet.max_row + 1):
        pole_match = bool(pole_column and target is not None and normalize_pole_number(worksheet.cell(row, pole_column).value) == target)
        internal_match = bool(internal_column and target_internal is not None and normalize_pole_number(worksheet.cell(row, internal_column).value) == target_internal)
        # Pole Number is the stable workbook-row key. Internal ID is the fallback for sheets
        # without a Pole Number column and remains admin-editable for backward compatibility.
        matched = pole_match if pole_column and target is not None else internal_match
        if matched: rows.append(row)
    return headers, rows, pole_column, internal_columns


def inspect_workbook(contents: bytes, pole_number, internal_id=None) -> dict:
    target = normalize_pole_number(pole_number)
    if target is None:
        raise ValueError("Selected pole has no Pole Number")
    try:
        workbook = load_workbook(BytesIO(contents), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError("Workbook is invalid or cannot be read") from exc
    poles_sheet, _ = detect_pole_sheet(workbook, require_location=False)
    _, pole_rows, pole_column, internal_columns = _matching_rows(poles_sheet, target, internal_id)
    if (pole_column == 0 and not internal_columns) or not pole_rows:
        raise ValueError(f"No exact pole identity row exists in worksheet {poles_sheet.title!r}")
    if len(pole_rows) > 1:
        raise ValueError(f"Workbook has duplicate Pole Number {target}")

    worksheets = []
    for worksheet in workbook.worksheets:
        headers, rows, pole_column, internal_columns = _matching_rows(worksheet, target, internal_id)
        if not rows:
            continue
        worksheets.append({
            "name": worksheet.title,
            "headers": [label for _, label in headers],
            "columns": [column for column, _ in headers],
            "pole_number_column": pole_column,
            "internal_id_columns": internal_columns,
            "identity_columns": ([pole_column] if pole_column else []) + internal_columns,
            "rows": [
                {
                    "row": row,
                    "values": [_serializable(worksheet.cell(row, column).value) for column, _ in headers],
                }
                for row in rows
            ],
        })
    return {"pole_number": target, "internal_id": target_internal if (target_internal := normalize_pole_number(internal_id)) else None,
            "pole_sheet": poles_sheet.title, "pole_row_count": len(pole_rows), "worksheets": worksheets}


def _coerce_value(value, existing):
    if value is None or value == "":
        return None
    if isinstance(value, (dict, list)):
        raise ValueError("Cell values must be scalar")
    if isinstance(value, str):
        if isinstance(existing, bool):
            normalized = value.strip().casefold()
            if normalized in {"true", "1", "yes"}:
                return True
            if normalized in {"false", "0", "no"}:
                return False
            raise ValueError("Enter true or false for this cell")
        if isinstance(existing, int) and not isinstance(existing, bool):
            try:
                return int(value.strip())
            except ValueError as exc:
                raise ValueError("Enter a whole number for this cell") from exc
        if isinstance(existing, float):
            try:
                number = float(value.strip())
            except ValueError as exc:
                raise ValueError("Enter a number for this cell") from exc
            if not math.isfinite(number):
                raise ValueError("Cell numbers must be finite")
            return number
        if isinstance(existing, datetime):
            try:
                return datetime.fromisoformat(value.strip())
            except ValueError as exc:
                raise ValueError("Enter an ISO date and time for this cell") from exc
        if isinstance(existing, date):
            try:
                return date.fromisoformat(value.strip())
            except ValueError as exc:
                raise ValueError("Enter an ISO date for this cell") from exc
        if isinstance(existing, time):
            try:
                return time.fromisoformat(value.strip())
            except ValueError as exc:
                raise ValueError("Enter an ISO time for this cell") from exc
        return value
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Cell numbers must be finite")
    return value


def apply_workbook_updates(contents: bytes, pole_number, updates: list[dict], *, internal_id=None, allow_internal_id: bool = False) -> bytes:
    if not updates:
        raise ValueError("No workbook changes were submitted")
    inspected = inspect_workbook(contents, pole_number, internal_id)
    allowed_rows = {
        (sheet["name"], row["row"]): {
            column: (header, sheet["pole_number_column"])
            for column, header in zip(sheet["columns"], sheet["headers"])
        }
        for sheet in inspected["worksheets"]
        for row in sheet["rows"]
    }
    internal_id_columns = {
        (sheet["name"], column) for sheet in inspected["worksheets"] for column in sheet["internal_id_columns"]
    }
    workbook = load_workbook(BytesIO(contents), data_only=False, read_only=False)
    touched = set()
    for update in updates:
        sheet_name, row_number, column_number = update.get("sheet"), update.get("row"), update.get("column")
        if not isinstance(sheet_name, str) or not isinstance(row_number, int) or isinstance(row_number, bool) or not isinstance(column_number, int) or isinstance(column_number, bool):
            raise ValueError("Workbook cell reference is invalid")
        row_columns = allowed_rows.get((sheet_name, row_number))
        if row_columns is None:
            raise ValueError("Workbook row does not belong to Pole Number")
        column_info = row_columns.get(column_number)
        if column_info is None:
            raise ValueError("Workbook column does not exist")
        _, pole_number_column = column_info
        if column_number == pole_number_column:
            raise ValueError("Pole Number is read-only")
        if not allow_internal_id and (sheet_name, column_number) in internal_id_columns:
            raise ValueError("internal_id can only be changed by an admin")
        marker = (sheet_name, row_number, column_number)
        if marker in touched:
            raise ValueError("A workbook cell was submitted more than once")
        touched.add(marker)
        cell = workbook[sheet_name].cell(row_number, column_number)
        cell.value = _coerce_value(update.get("value"), cell.value)
        if isinstance(cell.value, str):
            cell.data_type = "s"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
