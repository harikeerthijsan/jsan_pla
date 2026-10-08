from openpyxl import Workbook
import pytest
from app.ingest import parse_workbook

def test_dynamic_workbook_parser(tmp_path):
    p=tmp_path/'sample.xlsx'; wb=Workbook(); ws=wb.active; ws.title='poles'
    ws.append(['internal_id','pole_number','grade_of_construction','latitude','longitude','bottom_elev_ft','top_lat','top_lon','top_elev_ft','other_1_type','other_1_id','other_1_lat','other_1_lon','other_1_elev_ft','remarks'])
    ws.append([1,'P-100','A',34.37,-118.91,500,34.37,-118.91,540,'OTHER_POLE',1,34.37,-118.91,520,None])
    ws.append([2,'P-100','A',34.371,-118.911,501,34.371,-118.911,541,None,None,None,None,None,None])
    for name in ['attachments','crossarms','equipment','anchors','guys','sidewalk_braces','span_guys']:
        s=wb.create_sheet(name); s.append(['internal_id','pole_number'])
    wb.save(p)
    out=parse_workbook(str(p),'EPSG:6424')
    assert len(out['poles'])==2
    assert len([f for f in out['features'] if f['properties']['feature_type']=='pole'])==2
    assert any(f['rule_id']=='PLA-R006' and f['severity']=='FAIL' for f in out['findings'])
    dup=[f for f in out['findings'] if f['rule_id']=='PLA-R002']
    assert len(dup)==2


def test_generalized_catalogue_detects_aliases_xy_z_and_preserves_source_columns(tmp_path):
    path=tmp_path/'generalized.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='Customer Assets 2026'
    sheet.append(['Internal ID','Pole Number','Easting','Northing','Elevation','Customer Status','Unmapped Field'])
    sheet.append([101,'A-101',6450000.25,1840000.75,512.5,'Ready','keep exactly'])
    related=workbook.create_sheet('Custom Attachments')
    related.append(['InternalID','Asset Owner']); related.append([101,'Utility A'])
    workbook.save(path)

    result=parse_workbook(str(path),'EPSG:6424')
    assert result['pole_sheet']=='Customer Assets 2026'
    assert result['poles'][0]['internal_id']==101 and result['poles'][0]['pole_number']=='A-101'
    assert result['poles'][0]['source_x']==pytest.approx(6450000.25)
    assert result['poles'][0]['source_y']==pytest.approx(1840000.75)
    assert result['poles'][0]['source_z']==pytest.approx(512.5)
    assert result['poles'][0]['manifest']['Unmapped Field']=='keep exactly'
    assert not any(f['rule_id']=='PLA-R003' for f in result['findings'])
    location=next(f for f in result['features'] if f['properties']['feature_type']=='pole_location')
    assert location['geometry']['coordinates']==[6450000.25,1840000.75,512.5]


@pytest.mark.parametrize('headers,message',[
    (['Internal ID','Latitude','Longitude'],'Pole Number'),
    (['Internal ID','Pole Number','Latitude'],'latitude/longitude or X/Y'),
])
def test_generalized_catalogue_rejects_missing_required_fields(tmp_path,headers,message):
    path=tmp_path/'missing.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='Assets'
    sheet.append(headers); sheet.append([1,'P-1',34.2][:len(headers)])
    workbook.save(path)
    with pytest.raises(ValueError,match=message): parse_workbook(str(path),'EPSG:4326')


def test_generalized_catalogue_rejects_duplicate_internal_ids(tmp_path):
    path=tmp_path/'duplicates.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='Assets'
    sheet.append(['Internal ID','Pole Number','Latitude','Longitude'])
    sheet.append([1,'P-1',34.2,-118.1]); sheet.append([1,'P-2',34.3,-118.2]); workbook.save(path)
    with pytest.raises(ValueError,match='duplicate Internal ID 1'): parse_workbook(str(path),'EPSG:4326')


def test_height_and_pole_id_stay_ordinary_attributes():
    from app.source_schema import canonical_field
    assert canonical_field('Height') is None
    assert canonical_field('Pole ID') is None
    assert canonical_field('Elevation')=='z' and canonical_field('internal_id')=='internal_id'


def test_catalogue_with_height_column_is_not_treated_as_elevation(tmp_path):
    path=tmp_path/'height.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='Assets'
    sheet.append(['Internal ID','Pole Number','Pole ID','X','Y','Height'])
    sheet.append([7,'P-7','TAG-7',6450000.0,1840000.0,45])
    workbook.create_sheet('Empty notes')  # an entirely empty sheet must not crash detection
    workbook.save(path)
    result=parse_workbook(str(path),'EPSG:6424')
    pole=result['poles'][0]
    assert pole['source_z'] is None and pole['manifest']['Height']==45 and pole['manifest']['Pole ID']=='TAG-7'
    assert not any(f['properties']['feature_type'] in ('pole','pole_location') for f in result['features'])


def test_xy_catalogue_with_base_and_top_elevations_builds_pole_line_at_xy(tmp_path):
    path=tmp_path/'xy-elev.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='poles'
    sheet.append(['internal_id','pole_number','x','y','bottom_elev_ft','top_elev_ft'])
    sheet.append([3,'P-3',6450000.0,1840000.0,500.0,540.0]); workbook.save(path)
    result=parse_workbook(str(path),'EPSG:6424')
    pole=next(f for f in result['features'] if f['properties']['feature_type']=='pole')
    assert pole['geometry']['coordinates']==[[6450000.0,1840000.0,500.0],[6450000.0,1840000.0,540.0]]
    assert pole['properties']['source']=='workbook_xyz'
