"""PLA-1 (single ZoneName sheet): pole columns filled from the saved LiDAR points."""
import os
import uuid
from io import BytesIO

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.db import SessionLocal
from app.main import app
from app.models import Pole, ProductionAnnotation
from app.pla1_workbook import fill_pla1_workbook

HEADERS = ["internal_id", "Pole Number", "FQN #", "Latitude", "Longitude", "Grade of Construction", "First Pole", "Last Pole",
           "Top of Pole Elevation", "Bottom of Pole Elevation", "Incomplete Bottom", "Incomplete Latitude", "Incomplete Longitude",
           "1 Attachment Owner", "Remarks"]


def zone_name(rows):
    wb = Workbook(); ws = wb.active; ws.title = "ZoneName"; ws.append(HEADERS)
    for number in rows:
        ws.append([None, number, "FQN", 33.0, -118.0, "A", None, None, None, None, None, None, None, "Verizon", None])
    validation = DataValidation(type="list", formula1='"YES, Incomplete"'); validation.add(f"G2:H{len(rows) + 1}"); ws.add_data_validation(validation)
    out = BytesIO(); wb.save(out); return out.getvalue()


def point(pole, feature, *, lat=33.12345678912, lon=-118.98765432112, z=40.0, attrs=None, updated="2026-10-30T10:00:00+00:00"):
    return {"pole_internal_id": pole, "family": "poles", "feature_type": feature, "latitude": lat, "longitude": lon, "z": z,
            "attributes": attrs or {}, "updated_at": updated}


def cells(data, number):
    ws = load_workbook(BytesIO(data))["ZoneName"]
    names = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == number)
    return dict(zip(names, row)), ws


def test_pole_base_and_top_fill_the_pole_columns():
    points = [
        point(1, "Pole_Base", z=500.25, attrs={"first_pole": "INCOMPLETE", "last_pole": "YES",
                                               "prev_inc_lat": "33.2", "prev_inc_lon": "-118.3", "prev_inc_elev_ft": "501.5"}),
        point(1, "Pole_Top", z=538.75),
    ]
    data, summary = fill_pla1_workbook(zone_name(["P-1", "P-2"]), {1: "P-1", 2: "P-2"}, points)
    row, ws = cells(data, "P-1")
    assert row["Latitude"] == 33.12345679 and row["Longitude"] == -118.98765432
    assert row["Bottom of Pole Elevation"] == 500.25 and row["Top of Pole Elevation"] == 538.75
    assert row["First Pole"] == "Incomplete" and row["Last Pole"] == "YES"   # the Excel's own dropdown values
    assert row["Incomplete Latitude"] == 33.2 and row["Incomplete Longitude"] == -118.3 and row["Incomplete Bottom"] == 501.5
    # Everything else is left as it was.
    assert row["Grade of Construction"] == "A" and row["1 Attachment Owner"] == "Verizon" and row["FQN #"] == "FQN"
    untouched, _ = cells(data, "P-2")
    assert untouched["Latitude"] == 33.0 and untouched["Top of Pole Elevation"] is None
    assert summary["poles_filled"] == 1 and summary["skipped"] == [] and summary["missing_columns"] == []
    assert any(dv.formula1 == '"YES, Incomplete"' for dv in ws.data_validations.dataValidation)


def test_last_pole_incomplete_uses_the_next_pole_spot_and_metres_convert_to_feet():
    attrs = {"last_pole": "INCOMPLETE", "next_inc_lat": "33.4", "next_inc_lon": "-118.4", "next_inc_elev_ft": "160.1",
             "prev_inc_lat": "1", "prev_inc_lon": "1", "prev_inc_elev_ft": "1"}   # left over from an earlier first_pole choice
    data, _ = fill_pla1_workbook(zone_name(["P-1"]), {1: "P-1"}, [point(1, "Pole_Base", z=100.0, attrs=attrs)], elevation_to_ft=1 / 0.3048)
    row, _ = cells(data, "P-1")
    assert row["Bottom of Pole Elevation"] == 328.084 and row["Last Pole"] == "Incomplete" and row["First Pole"] is None
    assert row["Incomplete Latitude"] == 33.4 and row["Incomplete Longitude"] == -118.4 and row["Incomplete Bottom"] == 160.1


def test_latest_point_wins_and_unmatched_or_duplicated_poles_are_reported():
    points = [
        point(1, "Pole_Base", z=10.0, updated="2026-10-30T09:00:00+00:00"),
        point(1, "Pole_Base", z=12.0, updated="2026-10-30T11:00:00+00:00"),
        point(2, "Pole_Base", z=20.0),    # P-2 is on two rows: which one is the pole is unknown
        point(-7, "Pole_Base", z=30.0),   # a pole still without an Internal ID is matched by its Pole Number
        point(9, "Pole_Base", z=40.0),    # not in the Excel
    ]
    data, summary = fill_pla1_workbook(zone_name(["P-1", "P-2", "P-2", "P-7"]), {1: "P-1", 2: "P-2", -7: "P-7", 9: "P-9"}, points)
    assert cells(data, "P-1")[0]["Bottom of Pole Elevation"] == 12.0
    assert cells(data, "P-7")[0]["Bottom of Pole Elevation"] == 30.0
    ws = load_workbook(BytesIO(data))["ZoneName"]
    assert [r[9] for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == "P-2"] == [None, None]
    assert summary["poles_filled"] == 2 and sorted(summary["skipped"]) == ["P-2: on 2 rows of the Excel", "P-9: not in the Excel"]


def test_columns_the_excel_lacks_are_reported_not_added():
    wb = Workbook(); ws = wb.active; ws.title = "ZoneName"; ws.append(["Pole Number", "internal_id", "Latitude", "Longitude"]); ws.append(["P-1", None, 1, 2])
    raw = BytesIO(); wb.save(raw)
    data, summary = fill_pla1_workbook(raw.getvalue(), {1: "P-1"}, [point(1, "Pole_Top", z=50.0), point(1, "Pole_Base", z=10.0)])
    assert summary["missing_columns"] == ["bottom_elev", "top_elev"] and summary["poles_filled"] == 1
    assert [c.value for c in load_workbook(BytesIO(data))["ZoneName"][1]] == ["Pole Number", "internal_id", "Latitude", "Longitude"]


def test_pla1_download_is_filled_from_saved_points():
    from tests.test_workbook_format import dataset, login
    with TestClient(app) as client:
        headers = login(client)
        project_id = dataset(client, headers, zone_name(["P-1"]), "PLA-1")
        db = SessionLocal()
        try:
            db.add(Pole(project_id=project_id, internal_id=1, pole_number="P-1", manifest_json="{}"))
            db.add(ProductionAnnotation(id=uuid.uuid4().hex, project_id=project_id, block_name="tile", family="poles", feature_type="Pole_Top",
                                        x=1, y=2, z=538.5, latitude=33.5, longitude=-118.5, pole_internal_id=1, attributes_json="{}",
                                        created_by="admin@jsan.local", modified_by="admin@jsan.local"))
            db.commit()
        finally:
            db.close()
        latest = client.get(f"/api/projects/{project_id}/production-workbook/latest", headers=headers)
        assert latest.status_code == 200, latest.text
        assert latest.headers["x-pla1-summary"] == '{"poles_filled": 1, "skipped": 0, "missing_columns": []}'
        assert latest.headers["content-disposition"].endswith('client-updated.xlsx"') and latest.headers["x-workbook-edited"] == "0"
        assert cells(latest.content, "P-1")[0]["Top of Pole Elevation"] == 538.5
