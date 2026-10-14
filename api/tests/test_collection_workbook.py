"""Saved points → COLLECTION workbook: numbered column sets, dynamic sets, preserved Excel structure, API."""
import os
import re
import uuid
import zipfile
from io import BytesIO

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from app.collection_workbook import build_collection_workbook, restore_validation_extensions
from app.db import SessionLocal, initialize_schema
from app.main import app
from app.models import Pole, ProductionAnnotation, Project

X14 = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
       '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
       '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>lut_owners!$A$3:$A$40</xm:f></x14:formula1>'
       '<xm:sqref>{ref}</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')


def collection_template(rows=6):
    """A small workbook in the COLLECTION layout, like the customer template: tables, empty pre-made rows,
    a pole_number lookup formula, a standard dropdown and an x14 (lookup-list) dropdown on the owner columns."""
    wb = Workbook()
    poles = wb.active; poles.title = "poles"
    poles.append(["internal_id", "pole_number", "first_pole", "latitude", "longitude", "bottom_elev_ft", "top_lat", "top_lon", "top_elev_ft", "environment",
                  "other_1_type", "other_1_id", "other_1_attach_N", "other_1_lat", "other_1_lon", "other_1_elev_ft", "remarks"])
    poles.append([1, "P-1"]); poles.append([2, "P-2"])
    braces = wb.create_sheet("sidewalk_braces")
    braces.append(["internal_id", "pole_number"] + [f"swb_{n}_{s}" for n in (1, 2) for s in ("key", "owner", "anchor_no", "elev_ft", "length_ft", "guy_nos")] + ["remarks"])
    for r in range(2, rows + 2):
        braces.append([None, f'=IFERROR(INDEX(poles!$B:$B,MATCH($A{r},poles!$A:$A,0)),"")'])
    braces.add_table(Table(displayName="Braces", ref=f"A1:O{rows + 1}"))
    dv = DataValidation(type="list", formula1='"LOAD,SUPPORT"', allow_blank=True)
    dv.add(f"C2:C{rows + 1}"); dv.add(f"I2:I{rows + 1}")
    braces.add_data_validation(dv)
    wb.create_sheet("lut_owners").append(["owner"])
    raw = BytesIO(); wb.save(raw)
    # Add an x14 dropdown on the swb_n_owner columns (D and J), as Excel writes it.
    src = zipfile.ZipFile(BytesIO(raw.getvalue())); out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                data = data.decode().replace("</worksheet>", X14.replace("{ref}", f"D2:D{rows + 1} J2:J{rows + 1}") + "</worksheet>").encode()
            z.writestr(item, data)
    return out.getvalue()


def point(feature, family, attrs, pole=1, z=40.0):
    return {"pole_internal_id": pole, "family": family, "feature_type": feature, "latitude": 33.123456789, "longitude": -118.987654321,
            "z": z, "attributes": attrs, "updated_at": "2026-10-13T10:00:00+00:00"}


POLES = [{"internal_id": 1, "pole_number": "P-1"}, {"internal_id": 2, "pole_number": "P-2"}, {"internal_id": 7, "pole_number": "P-7"}]


def headers(ws):
    return [c.value for c in ws[1]]


def row(ws, internal_id):
    return next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 1).value) == str(internal_id))


def value(ws, internal_id, column):
    return ws.cell(row(ws, internal_id), headers(ws).index(column) + 1).value


def test_points_land_in_their_numbered_columns():
    data, summary = build_collection_workbook(collection_template(), POLES, [
        point("Pole_Base", "poles", {"first_pole": "YES", "environment": "Railroad", "other_type": "BUILDING", "other_id": "B-7"}, z=19.25),
        point("Pole_Top", "poles", {}, z=55.5),
        point("swb_1", "sidewalk_braces", {"swb_key": '2" Schedule 40', "swb_anchor_no": "anc_1", "swb_length_ft": "8.5", "swb_guy_nos": "guy_1"}),
    ])
    wb = load_workbook(BytesIO(data))
    poles, braces = wb["poles"], wb["sidewalk_braces"]
    assert summary["source"] == "workbook" and summary["points_written"] == 3 and summary["columns_added"] == []
    assert value(poles, 1, "latitude") == 33.12345679 and value(poles, 1, "bottom_elev_ft") == 19.25 and value(poles, 1, "top_elev_ft") == 55.5
    assert value(poles, 1, "first_pole") == "YES" and value(poles, 1, "environment") == "Railroad"
    assert value(poles, 1, "other_1_type") == "BUILDING" and value(poles, 1, "other_1_id") == "B-7"
    assert value(braces, 1, "swb_1_key") == '2" Schedule 40' and value(braces, 1, "swb_1_anchor_no") == "anc_1"
    assert value(braces, 1, "swb_1_length_ft") == "8.5" and value(braces, 1, "swb_1_elev_ft") == 40.0


def test_extra_point_number_adds_a_full_column_set_in_place():
    data, summary = build_collection_workbook(collection_template(), POLES, [point("swb_3", "sidewalk_braces", {"swb_owner": "Cox", "swb_anchor_no": "anc_2"})])
    wb = load_workbook(BytesIO(data)); braces = wb["sidewalk_braces"]; names = headers(braces)
    assert summary["columns_added"] == ["sidewalk_braces: swb_3_* (6 columns)"]
    assert names[-7:] == [f"swb_3_{s}" for s in ("key", "owner", "anchor_no", "elev_ft", "length_ft", "guy_nos")] + ["remarks"]
    assert value(braces, 1, "swb_3_owner") == "Cox" and value(braces, 1, "swb_3_anchor_no") == "anc_2"
    # The header looks like its neighbours and the table grows with the sheet.
    assert braces.cell(1, names.index("swb_3_key") + 1).font.b == braces.cell(1, names.index("swb_2_key") + 1).font.b
    table = list(braces.tables.values())[0]
    assert table.ref.startswith("A1:U") and [c.name for c in table.tableColumns] == names


def test_dropdowns_follow_and_extend_to_new_columns():
    data, _ = build_collection_workbook(collection_template(rows=6), POLES, [point("swb_3", "sidewalk_braces", {"swb_key": "x"})])
    wb = load_workbook(BytesIO(data)); braces = wb["sidewalk_braces"]
    key3 = braces.cell(1, headers(braces).index("swb_3_key") + 1).column_letter
    assert any(key3 in str(dv.sqref) for dv in braces.data_validations.dataValidation)       # standard dropdown copied
    xml = zipfile.ZipFile(BytesIO(data)).read("xl/worksheets/sheet2.xml").decode()
    owner3 = braces.cell(1, headers(braces).index("swb_3_owner") + 1).column_letter
    sqref = re.search(r"<xm:sqref>([^<]+)</xm:sqref>", xml).group(1).split()
    assert "lut_owners!$A$3:$A$40" in xml and f"{owner3}2:{owner3}7" in sqref and "D2:D7" in sqref  # x14 dropdown kept + copied


def test_poles_fill_premade_rows_and_keep_the_lookup_formula():
    data, summary = build_collection_workbook(collection_template(rows=6), POLES, [point("swb_1", "sidewalk_braces", {}, pole=7)])
    braces = load_workbook(BytesIO(data))["sidewalk_braces"]
    r = row(braces, 7)
    assert r <= 7 and braces.max_row == 7                         # used an empty template row, nothing appended
    assert braces.cell(r, 2).value.startswith("=IFERROR(INDEX(poles!$B:$B,MATCH($A") and f"$A{r}" in braces.cell(r, 2).value
    assert summary["rows_added"] >= 3                              # every pole is listed on every group sheet


def test_projects_without_a_collection_workbook_get_one_generated():
    data, summary = build_collection_workbook(None, POLES, [point("arm_2", "crossarms", {"arm_key": "Wood", "arm_owner": "SCE"})], elevation_to_ft=1 / 0.3048)
    wb = load_workbook(BytesIO(data))
    assert summary["source"] == "generated"
    assert {"poles", "attachments", "crossarms", "equipment", "anchors", "guys", "sidewalk_braces", "span_guys"} <= set(wb.sheetnames)
    arms = wb["crossarms"]
    assert value(arms, 1, "arm_2_key") == "Wood" and value(arms, 1, "arm_2_owner") == "SCE" and value(arms, 1, "arm_2_tip_lat") == 33.12345679
    assert abs(value(arms, 1, "arm_2_elev_ft") - 131.234) < 0.01   # 40 m in feet


def test_restoring_extensions_is_a_no_op_without_them():
    raw = BytesIO(); Workbook().save(raw)
    assert restore_validation_extensions(raw.getvalue(), raw.getvalue()) == raw.getvalue()


def test_download_route_needs_sign_in_and_returns_the_filled_workbook():
    initialize_schema()
    project_id = f"collection-{uuid.uuid4().hex[:8]}"
    db = SessionLocal()
    try:
        db.add(Project(id=project_id, name="Collection test", crs="EPSG:6424", units="US survey foot", status="READY"))
        db.add(Pole(project_id=project_id, internal_id=1, pole_number="P-1"))
        db.add(ProductionAnnotation(id=uuid.uuid4().hex, project_id=project_id, block_name="b", family="guys", feature_type="guy_9", pole_internal_id=1, x=1, y=2, z=30.5,
                                    latitude=33.1, longitude=-118.2, created_by="t", modified_by="t",
                                    attributes_json='{"guy_key": "1/2\\" EHS", "guy_anchor_no": "anc_1"}'))
        db.commit()
    finally:
        db.close()
    with TestClient(app) as client:
        assert client.get(f"/api/projects/{project_id}/production-workbook/collection").status_code == 401
        token = client.post("/api/auth/login", json={"email": "admin@jsan.local", "password": "ChangeMe123!"}).json()["token"]
        response = client.get(f"/api/projects/{project_id}/production-workbook/collection", headers={"Authorization": f"Bearer {token}"})
        assert response.status_code == 200, response.text
        assert response.headers["content-disposition"].endswith('-collection-updated.xlsx"')
        guys = load_workbook(BytesIO(response.content))["guys"]
        assert value(guys, 1, "guy_9_key") == '1/2" EHS' and value(guys, 1, "guy_9_anchor_no") == "anc_1" and value(guys, 1, "guy_9_elev_ft") == 30.5
        assert client.get("/api/projects/does-not-exist/production-workbook/collection", headers={"Authorization": f"Bearer {token}"}).status_code == 404


def test_arm_and_anchor_numbers_land_in_their_columns():
    data, _ = build_collection_workbook(None, POLES, [
        point("comm_1", "attachments_comm", {"comm_arm_no": "arm_2"}),
        point("util_3", "attachments_util", {"util_arm_no": "pole"}),
        point("guy_1", "guys", {"guy_anchor_no": "pole"}),
    ])
    wb = load_workbook(BytesIO(data))
    assert value(wb["attachments"], 1, "comm_1_arm_no") == "arm_2" and value(wb["attachments"], 1, "util_3_arm_no") == "pole"
    assert value(wb["guys"], 1, "guy_1_anchor_no") == "pole"


def test_poles_sheet_is_overwritten_with_lidar_values():
    # The template's own latitude/longitude (from the client) are replaced by the LiDAR-picked pole base and top.
    template = load_workbook(BytesIO(collection_template()))
    template["poles"]["D2"] = 34.0; template["poles"]["E2"] = -118.0
    raw = BytesIO(); template.save(raw)
    data, _ = build_collection_workbook(raw.getvalue(), POLES, [
        point("Pole_Base", "poles", {"environment": "Crossing Street"}, z=19.25),
        point("Pole_Top", "poles", {}, z=55.5),
    ])
    poles = load_workbook(BytesIO(data))["poles"]
    assert value(poles, 1, "latitude") == 33.12345679 and value(poles, 1, "longitude") == -118.98765432
    assert value(poles, 1, "bottom_elev_ft") == 19.25 and value(poles, 1, "top_elev_ft") == 55.5 and value(poles, 1, "top_lat") == 33.12345679
    assert value(poles, 1, "environment") == "Crossing Street"


def test_incomplete_first_and_last_pole_fill_prev_and_next_inc():
    data, summary = build_collection_workbook(None, POLES, [
        point("Pole_Base", "poles", {"first_pole": "INCOMPLETE", "prev_inc_lat": "33.1111", "prev_inc_lon": "-118.2222", "prev_inc_elev_ft": "40.5",
                                     "last_pole": "YES", "next_inc_lat": "33.9", "next_inc_lon": "-118.9", "next_inc_elev_ft": "41"}, z=10.0),
    ])
    poles = load_workbook(BytesIO(data))["poles"]
    assert value(poles, 1, "first_pole") == "INCOMPLETE"
    assert value(poles, 1, "prev_inc_lat") == 33.1111 and value(poles, 1, "prev_inc_lon") == -118.2222 and value(poles, 1, "prev_inc_elev_ft") == 40.5
    # last_pole is YES, so the next_inc values (left over from an earlier choice) are not written.
    assert value(poles, 1, "last_pole") == "YES" and value(poles, 1, "next_inc_lat") is None
    assert summary["points_written"] == 1


def test_children_on_a_parent_keep_their_own_lat_lon_and_on_the_pole_only_elevation():
    arm = {**point("arm_1", "crossarms", {"arm_key": "Wood"}, z=35.0), "latitude": 33.5, "longitude": -118.5}
    data, _ = build_collection_workbook(None, POLES, [
        arm,
        point("comm_1", "attachments_comm", {"parent_feature_id": "pole"}, z=22.0),       # on the pole
        point("comm_2", "attachments_comm", {"parent_feature_id": "arm_1"}, z=24.0),      # on arm_1
        point("eq_1", "equipment", {"parent_feature_id": "pole"}, z=28.0),
        point("util_1", "attachments_util", {"parent_feature_id": "arm_1"}, z=30.0),
    ])
    wb = load_workbook(BytesIO(data)); att, eq = wb["attachments"], wb["equipment"]
    assert value(att, 1, "comm_1_elev_ft") == 22.0 and value(att, 1, "comm_1_lat") is None and value(att, 1, "comm_1_lon") is None
    # On a parent: the child's own clicked position (not arm_1's 33.5 / -118.5).
    assert value(att, 1, "comm_2_lat") == 33.12345679 and value(att, 1, "comm_2_lon") == -118.98765432 and value(att, 1, "comm_2_elev_ft") == 24.0
    assert value(att, 1, "util_1_lat") == 33.12345679 and value(att, 1, "util_1_elev_ft") == 30.0
    assert value(eq, 1, "eq_1_elev_ft") == 28.0 and value(eq, 1, "eq_1_lat") is None


def test_child_on_the_pole_keeps_existing_lat_lon_cells():
    base = load_workbook(BytesIO(build_collection_workbook(None, POLES, [])[0]))
    att = base["attachments"]; names = headers(att)
    att.cell(row(att, 1), names.index("comm_1_lat") + 1, 34.1); att.cell(row(att, 1), names.index("comm_1_lon") + 1, -118.1)
    raw = BytesIO(); base.save(raw)
    data, _ = build_collection_workbook(raw.getvalue(), POLES, [point("comm_1", "attachments_comm", {"parent_feature_id": "pole"}, z=21.5)])
    att = load_workbook(BytesIO(data))["attachments"]
    assert value(att, 1, "comm_1_lat") == 34.1 and value(att, 1, "comm_1_lon") == -118.1 and value(att, 1, "comm_1_elev_ft") == 21.5


def test_other_poles_in_and_out_of_scope_fill_their_numbered_columns():
    attrs = {"other_1_type": "OTHER_POLE", "other_1_scope": "IN_SCOPE", "other_1_id": "P-77", "other_1_attach_N": "3",
             "other_2_type": "OTHER_POLE", "other_2_scope": "OUT_OF_SCOPE", "other_2_lat": "34.1234", "other_2_lon": "-118.4321", "other_2_elev_ft": "45.5",
             "other_3_type": "BUILDING",
             "other_5_type": "OTHER_POLE", "other_5_scope": "OUT_OF_SCOPE", "other_5_lat": "34.2", "other_5_lon": "-118.3", "other_5_elev_ft": "50"}
    data, summary = build_collection_workbook(None, POLES, [point("Pole_Base", "poles", attrs, z=10.0)])
    poles = load_workbook(BytesIO(data))["poles"]
    assert value(poles, 1, "other_1_type") == "OTHER_POLE" and value(poles, 1, "other_1_id") == "P-77" and value(poles, 1, "other_1_attach_N") == "3"
    assert value(poles, 1, "other_1_lat") is None                       # in scope: no position
    assert value(poles, 1, "other_2_lat") == 34.1234 and value(poles, 1, "other_2_lon") == -118.4321 and value(poles, 1, "other_2_elev_ft") == 45.5
    assert value(poles, 1, "other_2_id") is None                        # out of scope: no id
    assert value(poles, 1, "other_3_type") == "BUILDING"
    # The sheet has other_1…other_4; other_5 gets its own column set.
    assert value(poles, 1, "other_5_lat") == 34.2 and "poles: other_5_* (6 columns)" in summary["columns_added"]
    # The scope choice is not a workbook column.
    assert not any("scope" in str(h) for h in headers(poles))


def test_other_pole_saved_before_numbering_is_entry_one():
    data, _ = build_collection_workbook(None, POLES, [point("Pole_Base", "poles", {"other_type": "OTHER_POLE", "other_id": "P-5", "other_attach_N": "1"}, z=10.0)])
    poles = load_workbook(BytesIO(data))["poles"]
    assert value(poles, 1, "other_1_type") == "OTHER_POLE" and value(poles, 1, "other_1_id") == "P-5"
