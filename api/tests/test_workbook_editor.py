from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.workbook_editor import (
    apply_workbook_updates,
    inspect_workbook,
    match_geojson_pole_number,
)


def workbook_bytes():
    workbook = Workbook()
    poles = workbook.active
    poles.title = "poles"
    poles.append(["Internal ID", "Pole Number", "Remarks", "Height"])
    poles.append([1, "P-001", "North side", 12.5])
    poles.append([2, "P-002", "South side", 10])

    attachments = workbook.create_sheet("attachments")
    attachments.append(["Pole_Number", "Owner", "Usage"])
    attachments.append(["p-001", "Utility A", "Primary"])
    attachments.append(["P-001", "Utility B", None])
    attachments.append(["P-002", "Utility C", "Neutral"])

    unrelated = workbook.create_sheet("notes")
    unrelated.append(["Title", "Value"])
    unrelated.append(["Keep this", "unchanged"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_inspect_workbook_returns_all_rows_and_columns_for_pole_across_sheets():
    result = inspect_workbook(workbook_bytes(), "p-001")

    assert result["pole_row_count"] == 1
    sheets = {sheet["name"]: sheet for sheet in result["worksheets"]}
    assert set(sheets) == {"poles", "attachments"}
    assert sheets["poles"]["headers"] == ["Internal ID", "Pole Number", "Remarks", "Height"]
    assert sheets["poles"]["rows"] == [
        {"row": 2, "values": [1, "P-001", "North side", 12.5]}
    ]
    assert [row["row"] for row in sheets["attachments"]["rows"]] == [2, 3]
    assert sheets["attachments"]["rows"][1]["values"] == ["P-001", "Utility B", None]


def test_geojson_match_requires_exact_pole_number_property():
    features = [
        {"geometry_type": "Point", "properties": {"Pole Number": "p-001 "}},
        {"geometry_type": "Point", "properties": {"OBJECTID": 1}},
        {"geometry_type": "LineString", "properties": {"Pole Number": "P-001"}},
    ]

    assert match_geojson_pole_number(features, "P-001") == {
        "status": "MATCHED",
        "feature_indexes": [0],
    }
    assert match_geojson_pole_number(features[1:], "P-001")["status"] == "MISSING"
    assert match_geojson_pole_number(features[:1] * 2, "P-001")["status"] == "AMBIGUOUS"


def test_inspect_workbook_reports_duplicate_pole_number_as_ambiguous():
    data = workbook_bytes()
    workbook = load_workbook(BytesIO(data))
    workbook["poles"].append([3, "P-001", "Duplicate", 9])
    stream = BytesIO()
    workbook.save(stream)

    with pytest.raises(ValueError, match="duplicate Pole Number"):
        inspect_workbook(stream.getvalue(), "P-001")


def test_updates_preserve_source_bytes_and_unrelated_workbook_cells():
    source = workbook_bytes()
    updated = apply_workbook_updates(
        source,
        "P-001",
        [
            {"sheet": "poles", "row": 2, "column": 3, "value": "Updated"},
            {"sheet": "attachments", "row": 2, "column": 2, "value": "Utility Updated"},
            {"sheet": "attachments", "row": 3, "column": 3, "value": "Secondary"},
        ],
    )
    workbook = load_workbook(BytesIO(updated), data_only=False)

    assert source != updated
    assert workbook["poles"]["C2"].value == "Updated"
    assert workbook["attachments"]["B2"].value == "Utility Updated"
    assert workbook["attachments"]["C3"].value == "Secondary"
    assert workbook["notes"]["B2"].value == "unchanged"


def test_updates_cannot_change_pole_number_or_edit_an_unmatched_row():
    source = workbook_bytes()

    with pytest.raises(ValueError, match="Pole Number is read-only"):
        apply_workbook_updates(
            source,
            "P-001",
            [{"sheet": "poles", "row": 2, "column": 2, "value": "P-999"}],
        )

    with pytest.raises(ValueError, match="does not belong to Pole Number"):
        apply_workbook_updates(
            source,
            "P-001",
            [{"sheet": "poles", "row": 3, "column": 3, "value": "Wrong row"}],
        )


def test_formula_like_text_is_written_as_literal_text():
    updated = apply_workbook_updates(
        workbook_bytes(),
        "P-001",
        [{"sheet": "poles", "row": 2, "column": 3, "value": "=1+1"}],
    )
    workbook = load_workbook(BytesIO(updated), data_only=False)
    cell = workbook["poles"]["C2"]

    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_inspect_workbook_returns_safe_error_for_malformed_xlsx():
    with pytest.raises(ValueError, match="Workbook is invalid or cannot be read"):
        inspect_workbook(b"not an Excel workbook", "P-001")
