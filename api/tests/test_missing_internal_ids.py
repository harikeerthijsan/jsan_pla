"""Excel checked before upload; poles with a blank Internal ID load and get their ID later from the Production view."""
import os
import uuid
from io import BytesIO

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.auth import hash_password
from app.db import SessionLocal
from app.ingest import check_production_workbook, parse_workbook
from app.main import app
from app.models import AuditLog, Pole, ProductionAnnotation, User
from tests.test_workbook_format import dataset, login


def pole_workbook(rows, headers=("Internal ID", "Pole Number", "Latitude", "Longitude")):
    wb = Workbook(); ws = wb.active; ws.title = "poles"; ws.append(list(headers))
    for r in rows: ws.append(list(r))
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_checker_names_missing_columns_and_bad_rows():
    missing = check_production_workbook(pole_workbook([["P-1", 34.1]], headers=("Pole Number", "Latitude")))
    assert missing["errors"] == ["Sheet 'poles' is missing required columns: Internal ID, Latitude and Longitude (or X and Y)."]
    report = check_production_workbook(pole_workbook([[1, "P-1", 34.1, -118.2], [None, "P-2", 34.1, -118.2], ["abc", "P-3", 34, -118], [1, "P-4", 34, -118], [9, None, 34, -118]]))
    assert report["errors"] == ["Row 4: Internal ID 'abc' is not a whole number.", "Row 5: Internal ID 1 is also used on row 2.", "Row 6: Pole Number is empty (Internal ID 9)."]
    assert report["missing_internal_ids"] == 1 and "P-2" in report["warnings"][0] and report["pole_count"] == 5
    clean = check_production_workbook(pole_workbook([[1, "P-1", 34.1, -118.2], [None, "P-2", 34.1, -118.2]]))
    assert clean["errors"] == [] and clean["missing_internal_ids"] == 1


def test_parser_loads_blank_internal_ids_as_provisional(tmp_path):
    path = tmp_path / "poles.xlsx"; path.write_bytes(pole_workbook([[7, "P-7", 34.1, -118.2], [None, "P-8", 34.1, -118.2], [None, None, None, None]]))
    poles = {p["pole_number"]: p for p in parse_workbook(str(path), "EPSG:4326")["poles"]}
    assert poles["P-7"]["internal_id"] == 7 and not poles["P-7"]["provisional_internal_id"]
    assert poles["P-8"]["internal_id"] == -3 and poles["P-8"]["provisional_internal_id"]   # row 3, negative: never clashes
    bad = tmp_path / "bad.xlsx"; bad.write_bytes(pole_workbook([["x7", "P-7", 34.1, -118.2]]))
    with pytest.raises(ValueError, match="invalid Internal ID"):
        parse_workbook(str(bad), "EPSG:4326")


def test_processing_is_refused_before_conversion_when_columns_are_missing():
    with TestClient(app) as client:
        headers = login(client)
        project_id = dataset(client, headers, pole_workbook([["P-1", 34.1]], headers=("Pole Number", "Latitude")), None)
        refused = client.post(f"/api/projects/{project_id}/process-lidar", headers=headers)
        assert refused.status_code == 422 and "missing required columns: Internal ID" in refused.json()["detail"]
        check = client.post("/api/production/workbook-format", headers=headers, files={"file": ("c.xlsx", pole_workbook([[None, "P-2", 34.1, -118.2]]))})
        assert check.status_code == 200 and check.json()["check"]["missing_internal_ids"] == 1 and check.json()["check"]["errors"] == []


def make_user(role):
    email = f"setid-{role.lower()}-{uuid.uuid4().hex[:6]}@example.invalid"
    db = SessionLocal()
    try:
        db.add(User(email=email, name=f"{role} test", role=role, password_hash=hash_password("Unique-Test-Password-123!"))); db.commit()
    finally:
        db.close()
    return email


def test_setting_an_internal_id_moves_the_pole_and_writes_the_excel():
    user_email, qc_email = make_user("USER"), make_user("QC_REVIEWER")
    with TestClient(app) as client:
        admin = login(client)
        project_id = dataset(client, admin, pole_workbook([[5, "P-5", 34.1, -118.2], [None, "P-6", 34.1, -118.2]]), "PLA-1")
        db = SessionLocal()
        try:  # what the import creates: P-5 with its id, P-6 with a provisional one (row 3 -> -3), and some work on it
            db.add(Pole(project_id=project_id, internal_id=5, pole_number="P-5", manifest_json="{}"))
            db.add(Pole(project_id=project_id, internal_id=-3, pole_number="P-6", manifest_json="{}"))
            db.add(ProductionAnnotation(id=uuid.uuid4().hex, project_id=project_id, block_name="b", family="poles", feature_type="Pole_Base", pole_internal_id=-3,
                                        x=1, y=2, z=3, created_by="t", modified_by="t", attributes_json="{}"))
            db.commit()
        finally:
            db.close()
        url = f"/api/projects/{project_id}/poles/-3/internal-id"
        user = {"Authorization": "Bearer " + client.post("/api/auth/login", json={"email": user_email, "password": "Unique-Test-Password-123!"}).json()["token"]}
        qc = {"Authorization": "Bearer " + client.post("/api/auth/login", json={"email": qc_email, "password": "Unique-Test-Password-123!"}).json()["token"]}
        assert client.put(url, json={"internal_id": 66}).status_code == 401
        assert client.put(url, headers=qc, json={"internal_id": 66}).status_code == 403
        assert client.put(url, headers=user, json={"internal_id": 5}).status_code == 409          # already used by P-5
        assert client.put(url, headers=user, json={"internal_id": 0}).status_code == 422
        saved = client.put(url, headers=user, json={"internal_id": 66})                           # users may fill a missing id
        assert saved.status_code == 200 and saved.json() == {"internal_id": 66, "pole_number": "P-6", "workbook_updated": True}
        # Changing an existing id stays admin-only.
        assert client.put(f"/api/projects/{project_id}/poles/66/internal-id", headers=user, json={"internal_id": 67}).status_code == 403
        latest = client.get(f"/api/projects/{project_id}/production-workbook/latest", headers=admin)
        ws = load_workbook(BytesIO(latest.content))["poles"]
        assert [ws.cell(r, 1).value for r in range(2, 4)] == [5, 66] and latest.headers["x-workbook-edited"] == "1"
    db = SessionLocal()
    try:
        assert db.query(Pole).filter_by(project_id=project_id, pole_number="P-6").one().internal_id == 66
        assert db.query(ProductionAnnotation).filter_by(project_id=project_id).one().pole_internal_id == 66
        assert db.query(AuditLog).filter_by(action="SET_POLE_INTERNAL_ID", entity_id=f"{project_id}:66").count() == 1
    finally:
        db.close()


def test_collection_export_leaves_out_poles_without_an_id():
    import json
    with TestClient(app) as client:
        admin = login(client)
        project_id = dataset(client, admin, pole_workbook([[5, "P-5", 34.1, -118.2], [None, "P-6", 34.1, -118.2]]), "PLA-2")
        db = SessionLocal()
        try:
            db.add(Pole(project_id=project_id, internal_id=5, pole_number="P-5", manifest_json="{}"))
            db.add(Pole(project_id=project_id, internal_id=-3, pole_number="P-6", manifest_json="{}"))
            db.commit()
        finally:
            db.close()
        export = client.get(f"/api/projects/{project_id}/production-workbook/collection", headers=admin)
        assert export.status_code == 200 and json.loads(export.headers["x-collection-summary"])["poles_without_id"] == 1
        ids = [c.value for c in load_workbook(BytesIO(export.content))["poles"]["A"][1:]]
        assert 5 in ids and -3 not in ids


def test_pole_on_several_rows_gets_the_id_on_each_blank_row():
    from app.workbook_editor import set_pole_internal_id
    data = pole_workbook([[None, "P-6", 34.1, -118.2], [None, "P-6", 34.1, -118.2], [9, "P-6", 34.1, -118.2], [1, "P-1", 34.1, -118.2]])
    ws = load_workbook(BytesIO(set_pole_internal_id(data, "P-6", 66)))["poles"]
    # Both blank P-6 rows get 66; the P-6 row with another id (another pole) and other poles are untouched.
    assert [ws.cell(r, 1).value for r in range(2, 6)] == [66, 66, 9, 1]
    with pytest.raises(ValueError, match="no row"):
        set_pole_internal_id(pole_workbook([[9, "P-6", 34.1, -118.2]]), "P-6", 66)
