import json
import os
import uuid
from io import BytesIO

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import storage
from app.auth import hash_password
from app.db import SessionLocal, initialize_schema
from app.main import app
from app.models import AuditLog, DatasetFile, Pole, ProductionGeoFeature, Project
from app.workflow import DatasetVersion, VersionFile


def make_workbook():
    workbook = Workbook()
    poles = workbook.active
    poles.title = "poles"
    poles.append(["internal_id", "Pole Number", "remarks", "height"])
    poles.append([1, "P-001", "Original", 12.5])
    attachments = workbook.create_sheet("attachments")
    attachments.append(["Pole Number", "owner"])
    attachments.append(["P-001", "Utility A"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def auth_headers(client, email="admin@jsan.local", password="ChangeMe123!"):
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['token']}"}


def create_workbook_project(tmp_path, monkeypatch):
    initialize_schema()
    monkeypatch.setattr(storage, "LOCAL_ROOT", tmp_path)
    marker = uuid.uuid4().hex[:10]
    project_id = f"workbook-editor-{marker}"
    version_id = uuid.uuid4().hex
    source_id = uuid.uuid4().hex
    geojson_id = uuid.uuid4().hex
    source_key = f"{project_id}/versions/v1/source/COLLECTION.xlsx"
    geojson_key = f"{project_id}/versions/v1/source/poles.geojson"
    workbook_data = make_workbook()
    storage.upload_bytes(workbook_data, source_key, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    storage.upload_bytes(b"{}", geojson_key, "application/geo+json")

    db = SessionLocal()
    try:
        project = Project(id=project_id, name="Workbook editor test", customer="PLA", crs="EPSG:4326", units="m", status="LIDAR_READY")
        source = DatasetFile(id=source_id, project_id=project_id, filename="COLLECTION.xlsx", role="WORKBOOK", object_key=source_key, size_bytes=len(workbook_data), status="UPLOADED")
        geojson = DatasetFile(id=geojson_id, project_id=project_id, filename="poles.geojson", role="GEOJSON", object_key=geojson_key, size_bytes=2, status="UPLOADED")
        version = DatasetVersion(id=version_id, project_id=project_id, version_no=1, status="LIDAR_READY", created_by="admin@jsan.local")
        pole = Pole(project_id=project_id, internal_id=1, pole_number="P-001", manifest_json='{"pole_number":"P-001"}')
        feature = ProductionGeoFeature(
            project_id=project_id,
            source_file_id=geojson_id,
            feature_index=0,
            geometry_type="Point",
            source_crs="OGC:CRS84",
            source_geometry_json='{"type":"Point","coordinates":[10,20]}',
            geometry_json='{"type":"Point","coordinates":[10,20]}',
            x=10,
            y=20,
            properties_json='{"Pole Number":"p-001 "}',
        )
        db.add_all([project, source, geojson, version, pole, feature])
        db.flush()
        db.add_all([VersionFile(version_id=version_id, file_id=source_id), VersionFile(version_id=version_id, file_id=geojson_id)])
        qc_email = f"workbook-qc-{marker}@example.invalid"
        db.add(UserForTest(email=qc_email, name="Workbook QC", role="QC_REVIEWER", password_hash=hash_password("Unique-Test-Password-123!")))
        db.commit()
        return project_id, version_id, source_id, source_key, workbook_data, qc_email
    finally:
        db.close()


from app.models import User as UserForTest


def test_workbook_api_auth_save_snapshot_and_download(tmp_path, monkeypatch):
    project_id, version_id, source_id, source_key, original, qc_email = create_workbook_project(tmp_path, monkeypatch)

    with TestClient(app) as client:
        assert client.get(f"/api/projects/{project_id}/poles/1/workbook").status_code == 401
        admin = auth_headers(client)
        qc = auth_headers(client, qc_email, "Unique-Test-Password-123!")
        read = client.get(f"/api/projects/{project_id}/poles/1/workbook", headers=qc)
        assert read.status_code == 200, read.text
        assert read.json()["geojson_match"]["status"] == "MATCHED"
        assert read.json()["editable"] is False
        assert read.json()["download_available"] is False
        assert {sheet["name"] for sheet in read.json()["worksheets"]} == {"poles", "attachments"}
        assert read.json()["worksheets"][0]["rows"][0]["values"] == [1, "P-001", "Original", 12.5]

        forbidden = client.put(f"/api/projects/{project_id}/poles/1/workbook", headers=qc, json={
            "snapshot_file_id": source_id,
            "updates": [{"sheet": "poles", "row": 2, "column": 3, "value": "Denied"}],
        })
        assert forbidden.status_code == 403

        snapshot = client.get(f"/api/projects/{project_id}/poles/1/workbook", headers=admin).json()
        saved = client.put(f"/api/projects/{project_id}/poles/1/workbook", headers=admin, json={
            "snapshot_file_id": snapshot["snapshot_file_id"],
            "updates": [
                {"sheet": "poles", "row": 2, "column": 3, "value": "Edited"},
                {"sheet": "attachments", "row": 2, "column": 2, "value": "Utility Updated"},
            ],
        })
        assert saved.status_code == 200, saved.text
        assert saved.json()["editable"] is True
        assert saved.json()["download_available"] is True
        assert saved.json()["snapshot_file_id"] != source_id

        stale = client.put(f"/api/projects/{project_id}/poles/1/workbook", headers=admin, json={
            "snapshot_file_id": snapshot["snapshot_file_id"],
            "updates": [{"sheet": "poles", "row": 2, "column": 3, "value": "Stale"}],
        })
        assert stale.status_code == 409

        locked_key = client.put(f"/api/projects/{project_id}/poles/1/workbook", headers=admin, json={
            "snapshot_file_id": saved.json()["snapshot_file_id"],
            "updates": [{"sheet": "poles", "row": 2, "column": 2, "value": "P-999"}],
        })
        assert locked_key.status_code == 422

        second_save = client.put(f"/api/projects/{project_id}/poles/1/workbook", headers=admin, json={
            "snapshot_file_id": saved.json()["snapshot_file_id"],
            "updates": [{"sheet": "poles", "row": 2, "column": 3, "value": "Second edit"}],
        })
        assert second_save.status_code == 200, second_save.text
        assert second_save.json()["snapshot_file_id"] != saved.json()["snapshot_file_id"]

        downloaded = client.get(f"/api/projects/{project_id}/poles/1/workbook-download", headers=admin)
        assert downloaded.status_code == 200
        assert downloaded.headers["content-disposition"] == 'attachment; filename="COLLECTION-updated.xlsx"'
        result = load_workbook(BytesIO(downloaded.content), data_only=False)
        assert result["poles"]["C2"].value == "Second edit"
        assert result["attachments"]["B2"].value == "Utility Updated"
        assert storage.read_bytes(source_key) == original

        db = SessionLocal()
        try:
            saved_file = db.query(DatasetFile).filter_by(id=second_save.json()["snapshot_file_id"]).one()
            assert saved_file.role == "WORKBOOK_EDITED"
            assert db.query(VersionFile).filter_by(version_id=version_id, file_id=saved_file.id).one()
            audits = db.query(AuditLog).filter_by(action="UPDATE_PRODUCTION_WORKBOOK", entity_id=f"{project_id}:{version_id}:1").all()
            assert len(audits) == 2
            assert all("Edited" not in audit.detail_json and "Second edit" not in audit.detail_json for audit in audits)
            assert all("poles:remarks" in audit.detail_json for audit in audits)
        finally:
            db.close()
