"""Workbook format per Production dataset: PLA-1 (single ZoneName sheet, filled by hand) or PLA-2 (COLLECTION)."""
import os
import uuid
from io import BytesIO

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.collection_workbook import detect_workbook_format
from app.db import SessionLocal
from app.main import app
from app.models import Project


def zone_name_workbook():
    wb = Workbook(); ws = wb.active; ws.title = "ZoneName"
    ws.append(["Pole Number", "Latitude", "Longitude", "1 Attachment Owner"]); ws.append(["P-1", 33.1, -118.2, ""])
    out = BytesIO(); wb.save(out); return out.getvalue()


def collection_workbook():
    wb = Workbook(); wb.active.title = "poles"; wb["poles"].append(["internal_id", "pole_number"])
    for name in ("attachments", "crossarms", "anchors", "lut_owners"): wb.create_sheet(name)
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_detects_both_layouts():
    assert detect_workbook_format(zone_name_workbook()) == "PLA-1"
    assert detect_workbook_format(collection_workbook()) == "PLA-2"
    assert detect_workbook_format(b"not a workbook") is None


def login(client):
    return {"Authorization": "Bearer " + client.post("/api/auth/login", json={"email": "admin@jsan.local", "password": "ChangeMe123!"}).json()["token"]}


def upload(client, headers, project_id, name, role, data):
    prep = client.post(f"/api/projects/{project_id}/uploads/prepare", headers=headers, json={"filename": name, "role": role, "size_bytes": len(data)})
    assert prep.status_code == 200, prep.text
    put = client.put(prep.json()["upload_url"], headers=headers, files={"file": (name, data)})
    assert put.status_code == 200, put.text


def dataset(client, headers, workbook, fmt):
    body = {"name": f"Format test {uuid.uuid4().hex[:6]}", "crs": "EPSG:6424", "units": "US survey foot", **({"workbook_format": fmt} if fmt else {})}
    created = client.post("/api/projects", headers=headers, json=body)
    assert created.status_code == 200, created.text
    project_id = created.json()["id"]
    assert created.json()["workbook_format"] == fmt
    upload(client, headers, project_id, "client.xlsx", "WORKBOOK", workbook)
    upload(client, headers, project_id, "tile.laz", "LIDAR_SOURCE", b"LASF placeholder")
    return project_id


def test_mismatched_workbook_is_refused_before_processing():
    with TestClient(app) as client:
        headers = login(client)
        wrong = dataset(client, headers, zone_name_workbook(), "PLA-2")
        refused = client.post(f"/api/projects/{wrong}/process-lidar", headers=headers)
        assert refused.status_code == 422 and "PLA-1 (single ZoneName sheet)" in refused.json()["detail"]
        right = dataset(client, headers, zone_name_workbook(), "PLA-1")
        assert client.post(f"/api/projects/{right}/process-lidar", headers=headers).status_code == 200
        assert client.post("/api/projects", headers=headers, json={"name": "Bad format", "workbook_format": "PLA-3"}).status_code == 422


def test_projects_without_a_format_get_it_detected():
    with TestClient(app) as client:
        headers = login(client)
        legacy = dataset(client, headers, collection_workbook(), None)
        assert client.post(f"/api/projects/{legacy}/process-lidar", headers=headers).status_code == 200
    db = SessionLocal()
    try:
        assert db.query(Project).filter_by(id=legacy).first().workbook_format == "PLA-2"
    finally:
        db.close()


def test_pla1_downloads_the_latest_workbook_and_not_the_collection_export():
    with TestClient(app) as client:
        headers = login(client)
        project_id = dataset(client, headers, zone_name_workbook(), "PLA-1")
        assert client.get(f"/api/projects/{project_id}/production-workbook/latest").status_code == 401
        latest = client.get(f"/api/projects/{project_id}/production-workbook/latest", headers=headers)
        assert latest.status_code == 200
        assert latest.headers["content-disposition"].endswith('client-original.xlsx"') and latest.headers["x-workbook-edited"] == "0"
        assert detect_workbook_format(latest.content) == "PLA-1"
        collection = client.get(f"/api/projects/{project_id}/production-workbook/collection", headers=headers)
        assert collection.status_code == 409 and "Workbook Data" in collection.json()["detail"]
        assert client.get("/api/projects/missing-project/production-workbook/latest", headers=headers).status_code == 404


def test_format_check_endpoint_needs_upload_permission():
    from app.auth import hash_password
    from app.models import User
    email = f"format-qc-{uuid.uuid4().hex[:8]}@example.invalid"
    db = SessionLocal()
    try:
        db.add(User(email=email, name="QC only", role="QC_REVIEWER", password_hash=hash_password("Unique-Test-Password-123!"))); db.commit()
    finally:
        db.close()
    with TestClient(app) as client:
        url = "/api/production/workbook-format"
        assert client.post(url, files={"file": ("client.xlsx", zone_name_workbook())}).status_code == 401
        qc = {"Authorization": "Bearer " + client.post("/api/auth/login", json={"email": email, "password": "Unique-Test-Password-123!"}).json()["token"]}
        assert client.post(url, headers=qc, files={"file": ("client.xlsx", zone_name_workbook())}).status_code == 403
        headers = login(client)
        assert client.post(url, headers=headers, files={"file": ("client.xlsx", zone_name_workbook())}).json()["workbook_format"] == "PLA-1"
        assert client.post(url, headers=headers, files={"file": ("collection.xlsx", collection_workbook())}).json()["workbook_format"] == "PLA-2"
        assert client.post(url, headers=headers, files={"file": ("notes.csv", b"a,b")}).status_code == 422


def test_pla1_download_returns_the_latest_saved_edit():
    from app.models import DatasetFile
    from app.workflow import DatasetVersion, VersionFile
    from app.storage import upload_bytes
    with TestClient(app) as client:
        headers = login(client)
        project_id = dataset(client, headers, zone_name_workbook(), "PLA-1")
        db = SessionLocal()
        try:
            # What saving in the Workbook Data tab stores: a derived WORKBOOK_EDITED file on the active version.
            version = db.query(DatasetVersion).filter_by(project_id=project_id).order_by(DatasetVersion.version_no.desc()).first()
            edited = zone_name_workbook(); key = f"{project_id}/derived/edited.xlsx"; upload_bytes(edited, key, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            file_id = uuid.uuid4().hex
            db.add(DatasetFile(id=file_id, project_id=project_id, filename="client-updated.xlsx", role="WORKBOOK_EDITED", object_key=key, size_bytes=len(edited), status="UPLOADED"))
            db.flush(); db.add(VersionFile(version_id=version.id, file_id=file_id)); db.commit()
        finally:
            db.close()
        latest = client.get(f"/api/projects/{project_id}/production-workbook/latest", headers=headers)
        assert latest.status_code == 200 and latest.content == edited
        assert latest.headers["x-workbook-edited"] == "1" and latest.headers["content-disposition"].endswith('client-updated.xlsx"')
