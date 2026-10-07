import hashlib
import io
import json
import os
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_dynamic_api.db")
os.environ.setdefault("STORAGE_MODE", "local")
os.environ.setdefault("LOCAL_STORAGE_ROOT", "./test-storage")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ADMIN_EMAIL", "admin@jsan.local")
os.environ.setdefault("ADMIN_PASSWORD", "ChangeMe123!")

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import storage
from app.auth import hash_password
from app.db import SessionLocal, initialize_schema
from app.main import app
from app.models import AuditLog, DatasetFile, Finding, LidarBlock, Notification, Pole, ProcessingJob, Project, ReviewDecision, User
from app.notifications import notify_qc_finished
from app.workflow import DatasetVersion, VersionFile

PASSWORD = "Deliver-Test-Password-1!"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(label):
    workbook = Workbook()
    workbook.active.append(["internal_id", "Pole Number", "remarks"])
    workbook.active.append([1, "P-001", label])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _user(role="USER"):
    marker = uuid.uuid4().hex[:8]
    email = f"deliver-{role.lower()}-{marker}@example.invalid"
    db = SessionLocal()
    try:
        db.add(User(email=email, username=f"D{marker}", name=f"Deliver {marker}", role=role, password_hash=hash_password(PASSWORD)))
        db.commit()
        return email, db.query(User.id).filter_by(email=email).scalar()
    finally:
        db.close()


def _login(client, email, password=PASSWORD):
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['token']}"}


def _file(db, project_id, version_id, role, filename, data, created_at=None):
    file_id = uuid.uuid4().hex
    key = f"{project_id}/versions/v1/{role.lower()}/{file_id}-{filename}"
    storage.upload_bytes(data, key, XLSX)
    db.add(DatasetFile(id=file_id, project_id=project_id, filename=filename, role=role, object_key=key, size_bytes=len(data), status="UPLOADED",
                       created_at=created_at or datetime.now(timezone.utc)))
    db.flush()
    db.add(VersionFile(version_id=version_id, file_id=file_id))
    return key


def _datasets(tmp_path, monkeypatch):
    """Production dataset (edited Excel, 2 poles) and its QC dataset (v1 READY_FOR_QC, 2 findings incl. a formula-like message)."""
    initialize_schema()
    monkeypatch.setattr(storage, "LOCAL_ROOT", tmp_path)
    marker = uuid.uuid4().hex[:8]
    production_id, qc_id = f"deliver-{marker}", f"deliver-{marker}-qc"
    db = SessionLocal()
    try:
        db.add(Project(id=production_id, name=f"Deliver {marker}", customer="PLA", crs="EPSG:4326", units="degree", status="LIDAR_READY"))
        db.add(Project(id=qc_id, name=f"Deliver {marker} QC", customer="PLA", crs="EPSG:4326", units="degree", status="READY_FOR_QC", source_project_id=production_id))
        db.flush()
        db.add(ProcessingJob(id=uuid.uuid4().hex, project_id=production_id, job_type="LIDAR_INGEST", payload_json="{}", status="SUCCEEDED", progress=100, stage="done"))
        db.add(LidarBlock(project_id=production_id, name="tile-01", source_object_key=f"{production_id}/s.laz", object_key=f"{production_id}/t.copc.laz",
                          point_count=10, x_min=0, y_min=0, x_max=100, y_max=100, zmin=0, zmax=50, poles_json="[]"))
        p_version, q_version = uuid.uuid4().hex, uuid.uuid4().hex
        db.add(DatasetVersion(id=p_version, project_id=production_id, version_no=1, status="LIDAR_READY", created_by="admin@jsan.local"))
        db.add(DatasetVersion(id=q_version, project_id=qc_id, version_no=1, status="READY_FOR_QC", created_by="admin@jsan.local"))
        db.flush()
        _file(db, production_id, p_version, "WORKBOOK", "COLLECTION.xlsx", _xlsx("original"), datetime.now(timezone.utc) - timedelta(hours=2))
        _file(db, production_id, p_version, "WORKBOOK_EDITED", "COLLECTION-updated.xlsx", _xlsx("edited in production"), datetime.now(timezone.utc) - timedelta(hours=1))
        _file(db, qc_id, q_version, "WORKBOOK", "QC.xlsx", _xlsx("qc"))
        for project in (production_id, qc_id):
            for n in (1, 2):
                db.add(Pole(project_id=project, internal_id=n, pole_number=f"P-00{n}", manifest_json="{}", qc_status="FAIL" if n == 1 else "PASS", qc_fail=1 if n == 1 else 0))
        db.add(Finding(id=f"f1-{marker}", project_id=qc_id, internal_id=1, pole_number="P-001", rule_id="PLA-R010", severity="FAIL", sheet="poles",
                       field="guy", message='=HYPERLINK("http://evil.invalid","click")', related_poles_json="[]", status="OPEN"))
        db.add(Finding(id=f"f2-{marker}", project_id=qc_id, internal_id=1, pole_number="P-001", rule_id="PLA-R020", severity="REVIEW", sheet="poles",
                       field="height", message="Check height", related_poles_json="[]", status="OPEN"))
        db.flush()
        db.add(ReviewDecision(finding_id=f"f1-{marker}", reviewer_email="reviewer@example.invalid", decision="CONFIRMED_FAIL", comment="Confirmed"))
        db.commit()
    finally:
        db.close()
    return production_id, qc_id, q_version, marker


def _notifications(email):
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(email=email).one()
        return db.query(Notification).filter_by(user_id=user.id).order_by(Notification.id).all()
    finally:
        db.close()


def test_deliverable_package_is_frozen_at_approval_and_traceable(tmp_path, monkeypatch):
    production_id, qc_id, version_id, marker = _datasets(tmp_path, monkeypatch)
    user_email, _ = _user("USER")
    reviewer_email, _ = _user("QC_REVIEWER")
    with TestClient(app) as client:
        admin = _login(client, "admin@jsan.local", "ChangeMe123!")
        user = _login(client, user_email)
        client.post(f"/api/projects/{production_id}/production-annotations", headers=user,
                    json={"block_name": "tile-01", "family": "poles", "feature_type": "Pole_Base", "x": 10, "y": 20, "z": 5, "pole_internal_id": 1})
        url = f"/api/projects/{qc_id}/versions/{version_id}/deliverable"
        assert client.get(url).status_code == 401
        assert client.get(url, headers=_login(client, reviewer_email)).status_code == 403
        assert client.get(url, headers=user).status_code == 409

        assert client.post(f"/api/projects/{qc_id}/versions/{version_id}/approve", headers=admin).status_code == 200
        first = client.get(url, headers=user)
        assert first.status_code == 200, first.text
        data = first.json()
        assert data["built_at_approval"] is True and data["approved_by"] == "admin@jsan.local" and data["url"]
        paths = {f["path"] for f in data["files"]}
        assert paths == {"excel/COLLECTION-updated-production.xlsx", "excel/qc/QC.xlsx", "annotations/annotations.geojson", "reports/findings-report.xlsx"}

        db = SessionLocal()
        try:
            key = db.query(AuditLog).filter_by(action="BUILD_DELIVERABLE_PACKAGE", entity_id=version_id).one()
            package_key = json.loads(key.detail_json)["package_key"]
        finally:
            db.close()
        archive = zipfile.ZipFile(io.BytesIO(storage.read_bytes(package_key)))
        manifest = json.loads(archive.read("manifest.json"))
        assert {"README.txt", "manifest.json"} <= set(archive.namelist())
        for entry in manifest["files"]:
            assert hashlib.sha256(archive.read(entry["path"])).hexdigest() == entry["sha256"]
        production_excel = load_workbook(io.BytesIO(archive.read("excel/COLLECTION-updated-production.xlsx")))
        assert production_excel.active["C2"].value == "edited in production"
        geo = json.loads(archive.read("annotations/annotations.geojson"))
        assert geo["feature_count"] == 1 and geo["features"][0]["properties"]["point_name"] == "Pole_Base"

        report = load_workbook(io.BytesIO(archive.read("reports/findings-report.xlsx")))
        assert report.sheetnames == ["Summary", "Findings", "Corrections", "Poles"]
        findings = {row[1]: row for row in report["Findings"].iter_rows(min_row=2, values_only=True)}
        assert findings["PLA-R010"][11:14] == ("CONFIRMED_FAIL", "reviewer@example.invalid", findings["PLA-R010"][13])
        message_cell = next(row[7] for row in report["Findings"].iter_rows(min_row=2) if row[1].value == "PLA-R010")
        assert message_cell.value.startswith("=HYPERLINK") and message_cell.data_type == "s"

        # Later Production edits do not change what was delivered.
        db = SessionLocal()
        try:
            p_version = db.query(DatasetVersion).filter_by(project_id=production_id).one()
            _file(db, production_id, p_version.id, "WORKBOOK_EDITED", "COLLECTION-later.xlsx", _xlsx("after approval"))
            db.commit()
        finally:
            db.close()
        again = client.get(url, headers=user).json()
        assert again["package_sha256"] == data["package_sha256"] and again["generated_at"] == data["generated_at"]


def test_older_approved_versions_get_a_package_on_first_download(tmp_path, monkeypatch):
    _, qc_id, version_id, _ = _datasets(tmp_path, monkeypatch)
    db = SessionLocal()
    try:
        version = db.query(DatasetVersion).filter_by(id=version_id).one()
        # Approved between the original import (2 h ago) and the Production edit (1 h ago).
        version.status, version.approved_at = "APPROVED", datetime.now(timezone.utc) - timedelta(minutes=90)
        db.commit()
    finally:
        db.close()
    with TestClient(app) as client:
        admin = _login(client, "admin@jsan.local", "ChangeMe123!")
        data = client.get(f"/api/projects/{qc_id}/versions/{version_id}/deliverable", headers=admin).json()
        assert data["built_at_approval"] is False
        # The Production Excel is the one that existed at approval: the original import, not the later edit.
        paths = {f["path"] for f in data["files"]}
        assert "excel/COLLECTION-production.xlsx" in paths and "excel/COLLECTION-updated-production.xlsx" not in paths


def test_notifications_reach_the_right_people_without_self_or_admin_leaks(tmp_path, monkeypatch):
    production_id, qc_id, version_id, marker = _datasets(tmp_path, monkeypatch)
    owner_email, owner_id = _user("USER")
    reviewer_email, _ = _user("QC_REVIEWER")
    other_admin_email, _ = _user("ADMIN")
    with TestClient(app) as client:
        admin = _login(client, "admin@jsan.local", "ChangeMe123!")
        owner, reviewer = _login(client, owner_email), _login(client, reviewer_email)

        client.put(f"/api/projects/{production_id}/pole-assignments", headers=admin, json={"user_id": owner_id, "pole_internal_ids": [1, 2]})
        assigned = _notifications(owner_email)
        assert [n.kind for n in assigned] == ["POLES_ASSIGNED"] and assigned[0].title == "You were assigned 2 poles"
        assert not [n for n in _notifications("admin@jsan.local") if n.kind == "POLES_ASSIGNED"]

        # QC finishes: the pole owner hears about issues on their poles; admins get the summary.
        db = SessionLocal()
        try:
            notify_qc_finished(db, db.query(Project).filter_by(id=qc_id).one(), reviewer_email, True)
            db.commit()
        finally:
            db.close()
        issues = [n for n in _notifications(owner_email) if n.kind == "QC_ISSUES_ON_YOUR_POLES"]
        assert issues and issues[0].title == "QC found 2 issues on your poles" and "1 FAIL · 1 REVIEW on P-001" in issues[0].body
        assert json.loads(issues[0].link_json) == {"workspace": "PRODUCTION", "project_id": production_id, "pole_internal_id": 1}
        assert any(n.kind == "QC_FINISHED" for n in _notifications(other_admin_email))
        assert any(n.kind == "QC_FINISHED" for n in _notifications(reviewer_email))

        # Correction requested by the reviewer goes to the owner, not back to the reviewer.
        assert client.post(f"/api/findings/f1-{marker}/correction", headers=reviewer, json={"comment": "Fix the guy"}).status_code == 200
        requested = [n for n in _notifications(owner_email) if n.kind == "CORRECTION_REQUESTED"]
        assert requested and requested[0].title == "Correction requested on pole P-001" and "Fix the guy" in requested[0].body
        assert not [n for n in _notifications(reviewer_email) if n.kind == "CORRECTION_REQUESTED"]
        # Asking again for the same open correction does not notify twice.
        client.post(f"/api/findings/f1-{marker}/correction", headers=reviewer, json={"comment": "Again"})
        assert len([n for n in _notifications(owner_email) if n.kind == "CORRECTION_REQUESTED"]) == 1

        # An admin resolves it: the reviewer hears it is ready for re-check, without the admin's name.
        correction_id = client.get(f"/api/projects/{qc_id}/corrections", headers=admin).json()[0]["id"]
        assert client.post(f"/api/corrections/{correction_id}/resolve", headers=admin).status_code == 200
        resolved = [n for n in _notifications(reviewer_email) if n.kind == "CORRECTION_RESOLVED"]
        assert resolved and resolved[0].title.startswith("All corrections fixed") and "fixed by an admin" in resolved[0].body
        # Admins see the real name (the bootstrap admin has no username, so its display name).
        assert "fixed by JSAN QC Admin" in [n for n in _notifications(other_admin_email) if n.kind == "CORRECTION_RESOLVED"][0].body

        client.post(f"/api/projects/{qc_id}/versions/{version_id}/approve", headers=admin)
        assert any(n.kind == "VERSION_APPROVED" for n in _notifications(other_admin_email))
        assert not [n for n in _notifications("admin@jsan.local") if n.kind == "VERSION_APPROVED" and n.project_id == qc_id]


def test_notification_api_lists_and_marks_only_your_own():
    owner_email, owner_id = _user("USER")
    other_email, _ = _user("USER")
    db = SessionLocal()
    try:
        for title in ("one", "two"):
            db.add(Notification(user_id=owner_id, kind="TEST", title=title, link_json="{}"))
        db.commit()
    finally:
        db.close()
    with TestClient(app) as client:
        assert client.get("/api/notifications").status_code == 401
        owner, other = _login(client, owner_email), _login(client, other_email)
        listed = client.get("/api/notifications", headers=owner).json()
        assert listed["unread"] == 2 and [n["title"] for n in listed["items"]] == ["two", "one"]
        # Times carry an explicit UTC offset even on SQLite, so browsers do not read them as local time.
        assert all(n["created_at"].endswith("+00:00") for n in listed["items"])
        first_id = listed["items"][0]["id"]
        assert client.post(f"/api/notifications/{first_id}/read", headers=other).status_code == 404
        assert client.post(f"/api/notifications/{first_id}/read", headers=owner).json()["read"] is True
        assert client.get("/api/notifications", headers=owner).json()["unread"] == 1
        assert client.post("/api/notifications/read-all", headers=owner).json()["marked"] == 1
        assert client.get("/api/notifications", headers=owner).json()["unread"] == 0
        assert client.get("/api/notifications", headers=other).json()["unread"] == 0
